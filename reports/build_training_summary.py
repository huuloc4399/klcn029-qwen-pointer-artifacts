"""Create a measured-results report from pinned local artifacts, with provenance."""
from __future__ import annotations
import csv
import hashlib
import json
import math
import re
import statistics
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
STEM = 'BAO_CAO_TONG_HOP_DU_LIEU_HUAN_LUYEN_VA_DANH_GIA_2026-10-11'
SOURCES = {}

def sha(raw):
    return hashlib.sha256(raw).hexdigest()

def read(path):
    path = ROOT / path
    raw = path.read_bytes()
    SOURCES[path.relative_to(ROOT).as_posix()] = {'sha256':sha(raw),'bytes':len(raw)}
    return json.loads(raw.decode('utf-8-sig'))

def lines(path):
    path = ROOT / path
    raw = path.read_bytes()
    SOURCES[path.relative_to(ROOT).as_posix()] = {'sha256':sha(raw),'bytes':len(raw)}
    return [json.loads(line) for line in raw.decode('utf-8-sig').splitlines() if line.strip()]

def bind_zip(path):
    path = ROOT / path
    SOURCES[path.relative_to(ROOT).as_posix()] = {'sha256':sha(path.read_bytes()),'bytes':path.stat().st_size}
    return zipfile.ZipFile(path)

def zipped_json(archive, suffix):
    matches = [n for n in archive.namelist() if n.endswith(suffix)]
    if len(matches) != 1:
        raise ValueError('Ambiguous member: '+suffix)
    return json.loads(archive.read(matches[0]))

def percentile(values, q):
    values = sorted(values)
    p = (len(values)-1)*q
    low = int(p)
    return values[low]+(values[min(low+1,len(values)-1)]-values[low])*(p-low)

def runtime(rows):
    times = [r['runtime_seconds'] for r in rows]
    return {'count':len(rows),'input_tokens_total':sum(r.get('input_tokens',0) for r in rows), 'input_tokens_available':all('input_tokens' in r for r in rows), 'generated_tokens_total':sum(r['generated_tokens'] for r in rows),'generated_tokens_mean':statistics.mean(r['generated_tokens'] for r in rows),'runtime_seconds_total':sum(times),'runtime_seconds_mean':statistics.mean(times),'runtime_seconds_p50':percentile(times,.5),'runtime_seconds_p95':percentile(times,.95),'ended_with_eos_count':sum(r.get('ended_with_eos') is True for r in rows),'hit_max_new_tokens_count':sum(bool(r.get('hit_max_new_tokens')) for r in rows)}

def raw_rows(path, adapter_validation=False):
    with bind_zip(path) as archive:
        suffix = 'predictions.jsonl' if adapter_validation else 'predictions_raw.jsonl'
        member = next(n for n in archive.namelist() if n.endswith(suffix))
        rows = [json.loads(line) for line in archive.read(member).decode().splitlines()]
        if rows and 'report' in rows[0]:
            rows = [r['report'] for r in rows]
        return rows

def training(path):
    with bind_zip(path) as archive:
        if archive.testzip() is not None:
            raise ValueError('CRC error: '+path)
        manifest = zipped_json(archive,'/_FILE_MANIFEST.json')
        root = next(n for n in archive.namelist() if n.endswith('/_FILE_MANIFEST.json')).rsplit('/',1)[0]
        for name, expected in manifest['files'].items():
            raw = archive.read(root+'/'+name)
            if sha(raw) != expected['sha256'] or len(raw) != expected['size']:
                raise ValueError('Adapter member changed: '+name)
        raw = archive.read(root+'/adapter_model.safetensors')
        header_len = int.from_bytes(raw[:8],'little')
        header = json.loads(raw[8:8+header_len])
        tensors = [v for k,v in header.items() if k!='__metadata__']
        state = zipped_json(archive,'/trainer_state.json')
        return {'artifact':path,'manifest_verified_members':len(manifest['files']),'contract':zipped_json(archive,'/run_config.json'),'metrics':zipped_json(archive,'/metrics.json'),'adapter_config':zipped_json(archive,'/adapter_config.json'),'trainable_parameter_count_from_saved_tensors':sum(math.prod(t['shape']) for t in tensors),'saved_tensor_count':len(tensors),'saved_dtypes':dict(Counter(t['dtype'] for t in tensors)),'weights_sha256':sha(raw),'evaluation_history':[r for r in state['log_history'] if 'eval_loss' in r],'training_log_history':state['log_history'],'final_step':state['global_step'],'epoch':state['epoch']}

def collect():
    gold_path = 'training/evaluation/results/gold40_b0_b1_p0_p1_v1/'
    manifest = read(gold_path+'RESULT_MANIFEST.json')
    for name, expected in manifest['files'].items():
        if sha((ROOT/gold_path/name).read_bytes()) != expected:
            raise ValueError('Frozen GOLD40 result changed: '+name)
    data = {'date':'2026-10-11','scope':'OCR and CV extraction only','runs':{
        'B1':training('training/final_adapter-full_1.zip'),
        'P0_smoke':training('training/qwen3_4b_cvschema2_pointer_stage2_smoke_v1_final_adapter.zip'),
        'P0_full':training('training/qwen3_4b_cvschema2_pointer_stage2_full_v1_final_adapter.zip')},
        'dataset_manifest':read('training/data/resume_parsing_vision_cvschema_v2/manifest.json'),
        'fulljson_conversion':read('training/data/resume_parsing_vision_cvschema_v2/conversion_report.json'),
        'pointer_conversion':read('training/data/resume_parsing_vision_cvschema_v2_pointer_v1/conversion_report.json'),
        'generalization':read('training/audits/gold40_generalization_2026-10-11_v2/diagnostics.json'),
        'gold40':read(gold_path+'summary.json'),
        'real80_p0':read('training/evaluation/results/pointer_real80_p0_v1/evaluation_summary.json'),
        'real80_p1':read('training/evaluation/results/pointer_real80_selective_retry_p1_v1/evaluation_summary.json'),
        'ocr_paddle_v5':read('experiments/ocr/paddleocr_pilot_v1/summary.json'),
        'ocr_paddle_v3':read('experiments/ocr/paddleocr_pilot_v3_rec_v1/summary.json'),
        'ocr_hybrid':read('experiments/ocr/vietocr_hybrid_pilot_v1/comparison_summary.json'),
        'synthetic':{}}
    for name, expected in read('training/audits/gold40_generalization_2026-10-11_v2/MANIFEST.json')['files'].items():
        if sha((ROOT/'training/audits/gold40_generalization_2026-10-11_v2'/name).read_bytes())!=expected:
            raise ValueError('Diagnostics changed')
    for name,expected in data['dataset_manifest']['files_sha256'].items():
        if sha((ROOT/'training/data/resume_parsing_vision_cvschema_v2'/name).read_bytes())!=expected:
            raise ValueError('Dataset changed: '+name)
    for split, folder, paths in (
        ('validation','validation_base_vs_adapter_v1',{'B0':'training/validation-27-09-2026_greedy_v1.zip','B1':'training/validation-27-09-2026.zip'}),
        ('test','locked_test_base_vs_adapter_v1',{'B0':'training/qwen3_4b_base_cvschema2_test_raw_greedy_v1.zip','B1':'training/qwen3_4b_adapter_cvschema2_test_raw_greedy_v1.zip'})):
        comp = read('training/evaluation/results/'+folder+'/comparison_summary.json')
        pointer_folder = 'pointer_stage2_validation_p0_v1' if split=='validation' else 'pointer_stage2_locked_test_p0_v1'
        pointer = read('training/evaluation/results/'+pointer_folder+'/evaluation_summary.json')
        efficiency = {m:runtime(raw_rows(p,adapter_validation=(split=='validation' and m=='B1'))) for m,p in paths.items()}
        data['synthetic'][split]={'fulljson':comp,'pointer':pointer,'fulljson_efficiency':efficiency}
    with bind_zip('training/validation-27-09-2026.zip') as archive:
        data['historical_validation_before_frozen_parser']=zipped_json(archive,'summary.json')
    gold_manifest = read('annotation/audits/gold40_final_2026-10-10_complete/gold40_manifest.json')
    for cv_id, row in gold_manifest['exports'].items():
        for suffix,key in (('.json','gold_sha256'),('.status.json','status_sha256')):
            if sha((ROOT/f'annotation/gold_v2/{cv_id}{suffix}').read_bytes())!=row[key]:
                raise ValueError('Gold changed: '+cv_id)
    inputs = lines('experiments/baseline80/real_cv_pointer_inputs_v1/all_80_pointer_inputs.jsonl')
    audit_path=ROOT/'baseline_extraction/data/input_audit.csv'
    SOURCES[audit_path.relative_to(ROOT).as_posix()]={'sha256':sha(audit_path.read_bytes()),'bytes':audit_path.stat().st_size}
    audit=list(csv.DictReader(audit_path.open(encoding='utf-8-sig')))
    gold_ids=set(gold_manifest['exports'])
    data['real_data']={'pdf_count':len(audit),'page_count':sum(int(r['pages']) for r in audit),'languages':dict(Counter(r['language'] for r in audit)),'routes':dict(Counter(r['source_route'] for r in inputs)),'canonical_ids_in_frozen_80_inputs':len({r['canonical_id'] for r in inputs}),'gold_cv_count':len(gold_ids),'gold_pages':sum(int(r['pages']) for r in audit if r['id'] in gold_ids),'gold_canonical_ids':len({r['canonical_id'] for r in inputs if r['sample_id'] in gold_ids}),'gold_manifest':gold_manifest}
    data['training_data_pages']={}
    for split in ('train','validation','test'):
        rows=lines(f'training/data/resume_parsing_vision_cvschema_v2/{split}.jsonl')
        data['training_data_pages'][split]=dict(Counter(str(r['metadata']['num_pages']) for r in rows))
    data['sources']=SOURCES
    return data

def fmt(value, digits=2):
    if value is None:
        return 'Không áp dụng'
    return f'{value:,.{digits}f}'.replace(',','@').replace('.',',').replace('@','.')

def pct(value, digits=2):
    return fmt(value*100,digits)+'%' if value is not None else 'Không áp dụng'

def frac(n,total):
    return f'{n}/{total} ({pct(n/total)})'

def report(data):
    blocks=[]
    def h(text,level=1): blocks.append(('heading',level,text))
    def p(text): blocks.append(('paragraph',text))
    def table(headers,rows,widths=None): blocks.append(('table',headers,rows,widths))
    def bullet(text): blocks.append(('bullet',text))
    h('Báo cáo dữ liệu huấn luyện và đánh giá mô đun trích xuất CV',0)
    p('Đề tài CNTT KLCN029 | Báo cáo cho giảng viên hướng dẫn | Cập nhật ngày 11 tháng 10 năm 2026')
    p('Nhóm đã hoàn tất hai giai đoạn QLoRA trên Qwen3 4B, đánh giá sinh tự do trên validation và test tổng hợp, chạy Pointer trên 80 CV thật và so sánh bốn cấu hình trên GOLD40. Fine-tuning cải thiện rõ trên test tổng hợp, nhưng chưa vượt Qwen gốc trên CV thật. Kiểm chứng và retry tăng tính hợp lệ đầu ra Pointer, đồng thời phát sinh thêm lần gọi model.')
    p('Phạm vi báo cáo là dữ liệu, OCR và trích xuất CVSchema 2.0. Các số dưới đây không đo chất lượng matching CV–JD, chấm điểm hoặc sinh gợi ý của các mô đun khác. Các bảng đều gắn với artifact đã lưu; các thí nghiệm mới sau rà gold chưa được thực hiện.')
    h('1 Dữ liệu và cách sử dụng')
    real=data['real_data']; manifest=data['dataset_manifest']; diag=data['generalization']
    table(['Bộ dữ liệu','Số CV','Ngôn ngữ','Vai trò'],[
        ['Train tổng hợp','850','850 EN','Cập nhật gradient cả stage 1 và stage 2'],
        ['Validation tổng hợp','75','75 EN','Theo dõi loss và phân tích lỗi'],
        ['Test tổng hợp','75','75 EN','Đánh giá sau khóa parser'],
        ['CV thật','80 / 110 trang','60 EN và 20 VI','Đo cấu trúc và hiệu năng P0/P1'],
        ['GOLD40 thuộc real80','40 / 60 trang','20 EN và 20 VI','Đánh giá nội dung hậu kiểm B0/B1/P0/P1']], [2,1,1.25,2.75])
    p('Dataset tổng hợp là sukhrobnurali/resume-parsing-vision, revision '+manifest['source']['revision']+', giấy phép CC-BY-4.0 theo manifest nguồn. Bộ đã được chuyển thành CVSchema 2.0 và Hybrid Pointer v1. Hai stage dùng lại cùng 850 CV train, không phải 1.700 CV khác nhau. Chỉ split train dùng cập nhật gradient.')
    p('Dữ liệu tổng hợp không có source/target trùng hoàn toàn giữa train, validation và test theo audit đã đo. Ba template train gồm sidebar 271, two_column 267 và classic 312. Không có trùng nguyên văn không bảo đảm độc lập về mẫu sinh hoặc cấu trúc gần trùng.')
    p(f'Input đóng băng real80 có {real["canonical_ids_in_frozen_80_inputs"]} canonical ID trên 80 bản ghi. GOLD40 có {real["gold_canonical_ids"]} canonical ID cho 40 bản ghi theo mapping đó. Bộ corpus ở giai đoạn chuẩn bị khác có quy tắc deduplicate khác; không dùng số lượng của một phiên bản thay cho số mẫu của benchmark đã khóa.')
    p('GOLD40 đã được tổng hợp từ hai thành viên và phân xử; nhãn A/B trước phân xử không còn lưu nên không thể đo IAA lịch sử. Ba mẫu VI003, VI005, VI017 từng được trợ lý hỗ trợ hoàn tất theo yêu cầu. Chất lượng nội dung đang được rà lại bằng PDF; chưa có gold v2 được chấp nhận.')
    h('2 Phương pháp và cấu hình đã chạy')
    table(['Mã','Model và biểu diễn','Vai trò'],[
        ['B0','Qwen3 4B gốc, Full JSON','Baseline zero-shot'],
        ['B1','Adapter stage 1, Full JSON','Fine-tuning trên nhãn tổng hợp'],
        ['P0','Adapter stage 2, Hybrid Pointer','Sinh trường ngắn và khoảng dòng cho trường dài'],
        ['P1','P0 và validator với một retry tối đa','Chọn CV lỗi và sinh lại toàn JSON Pointer']], [0.5,2.75,3.75])
    p('Model chung là Qwen/Qwen3-4B-Instruct-2507, revision 1b4199c4f36b0cef378bfb12390c18780c18af4c. Stage 2 tiếp tục adapter stage 1 trên nhãn Pointer; full stage 2 không dùng adapter smoke làm điểm xuất phát. B0 và B1 dùng cùng prompt Full JSON trong từng benchmark; Pointer thêm đánh số dòng và hợp đồng span.')
    p('Nhóm chọn QLoRA để thực hiện yêu cầu có huấn luyện trong giới hạn Colab T4; kết quả VRAM bên dưới kiểm chứng khả năng chạy cấu hình này. Full fine-tuning cần cập nhật toàn bộ base, còn prompt-only là baseline B0 và không đáp ứng riêng yêu cầu huấn luyện. Các siêu tham số là cấu hình đã chốt của run, chưa có tìm kiếm siêu tham số để chứng minh tối ưu.')
    p('Hybrid Pointer dùng khoảng dòng cho summary, experience.description và projects.details nhằm bảo toàn nguyên văn và giảm phần văn bản model phải sinh. Lựa chọn một lượt giữ các trường của cùng entity trong một object và giảm lặp input so với nhiều prompt chuyên biệt; đề tài chưa có đối chứng nhiều prompt để kết luận phương án một lượt tốt hơn. P1 bổ sung kiểm chứng/retry cho lỗi cấu trúc và span có thể phát hiện mà không cần gold; chưa sửa được mọi lỗi ngữ nghĩa.')
    c1=data['runs']['B1']['contract']['training_contract']; c2=data['runs']['P0_full']['contract']['training_contract']
    table(['Thông số','Stage 1 Full JSON','Stage 2 Pointer full'],[
        ['Train / validation','850 / 75','850 / 75'],['Epoch / optimizer steps','1 / 107','1 / 107'],
        ['Batch mỗi thiết bị / accumulation','1 / 8','1 / 8'],['Batch hiệu dụng','8','8'],
        ['Learning rate',str(c1['learning_rate']),str(c2['learning_rate'])],
        ['Optimizer / scheduler','paged_adamw_8bit / cosine','paged_adamw_8bit / cosine'],
        ['Warmup ghi trong contract','warmup_steps = 0.03','warmup_steps = 0.03'],
        ['Weight decay / max grad norm','0.01 / 0.3','0.01 / 0.3'],
        ['LoRA rank / alpha / dropout','16 / 32 / 0.05','16 / 32 / 0.05'],
        ['LoRA target modules','all-linear','all-linear'],
        ['Lượng tử hóa base','4-bit NF4, double quantization','4-bit NF4, double quantization'],
        ['Precision / attention','FP16 / SDPA','FP16 / SDPA'],
        ['Maximum sequence length','5.120 token','5.120 token'],
        ['Loss / masking','chunked_nll / completion-only','chunked_nll / completion-only'],
        ['Packing / seed','False / 42','False / 42'],
        ['Lưu checkpoint / eval','Mỗi 20 / 40 bước','Mỗi 20 / 40 bước']], [2.6,2.2,2.2])
    p('Cấu hình warmup được chép nguyên khóa và giá trị trong contract để tái lập; không tự đổi thành số bước nguyên hoặc phần trăm khác. Gradient checkpointing được bật. Checkpoint có optimizer state để resume; local giữ 2 checkpoint, Drive giữ 3 theo contract. Tham số adapter trainable được nâng FP32 trong luồng FP16 đã sửa lỗi AMP.')
    r=data['runs']['B1']; p(f'Adapter lưu {r["saved_tensor_count"]} tensor FP32 với {r["trainable_parameter_count_from_saved_tensors"]:,} tham số LoRA, tương đương khoảng 33,03 triệu. Base được lượng tử hóa và giữ cố định; đây là QLoRA, không phải full fine-tuning tất cả tham số base.')
    p('Các run huấn luyện dùng Colab T4 theo hồ sơ notebook/audit. Môi trường B0/B1 GOLD40 được ghi trực tiếp là Tesla T4, Python 3.13.15, PyTorch 2.11.0+cu130 và CUDA 13.0. Đây là môi trường benchmark GOLD40, không thay thế log môi trường của mọi run trước.')
    pins=c1['software_pins']; table(['Thư viện trong contract train','Phiên bản'],list(pins.items()),[4.7,2.3])
    h('3 Kết quả huấn luyện đã đo')
    rows=[]
    for name,key in [('Stage 1 full','B1'),('Stage 2 smoke','P0_smoke'),('Stage 2 full','P0_full')]:
        run=data['runs'][key]; m=run['metrics']; c=run['contract']['training_contract']
        rows.append([name,f'{c["train_examples"]}/{c["validation_examples"]}',str(m['final_step']),fmt(m['train_loss'],6),fmt(m['final_eval_loss'],6),pct(m['final_eval_mean_token_accuracy'],4)])
    table(['Run','Train/val','Steps','Train loss','Eval loss','Token accuracy'],rows,[1.45,.8,.55,1.3,1.3,1.6])
    rows=[]
    for name,key in [('Stage 1 full','B1'),('Stage 2 smoke','P0_smoke'),('Stage 2 full','P0_full')]:
        m=data['runs'][key]['metrics']; rows.append([name,fmt(m['train_runtime']),fmt(m['train_runtime']/60),fmt(m['final_eval_runtime']),fmt(m['peak_gpu_memory_allocated_gb'],3),fmt(m['peak_gpu_memory_reserved_gb'],3)])
    table(['Run','Trainer giây','Trainer phút','Eval cuối giây','Allocated GiB','Reserved GiB'],rows,[1.45,1.1,1.1,1.15,1.1,1.1])
    p('VRAM dùng GiB vì mã đo chia byte cho 1024³, dù khóa JSON có hậu tố _gb. Reserved và allocated là hai đại lượng khác nhau, không cộng lại. Trainer runtime và runtime phiên hoàn tất đã ghi không bao gồm toàn bộ cài đặt, tải model, thời gian chờ hoặc mất kết nối; stage 1 có resume nên 38,21 phút không được coi là tổng thời gian từ đầu đến cuối hoặc dùng để kết luận train nhanh hơn stage 2.')
    p('Smoke dùng 16 mẫu train và 8 mẫu validation dài nhất, chỉ 2 optimizer steps. Nó kiểm tra khả năng chạy, lưu trọng số và parse output; không đại diện chất lượng mô hình full. Mẫu quick smoke có hai span chọn nhầm dòng tên dự án thay vì mô tả.')
    h('3 1 Loss validation theo các mốc đã lưu',2)
    rows=[]
    for label,key in [('Stage 1','B1'),('Stage 2 full','P0_full')]:
        for row in data['runs'][key]['evaluation_history']:
            rows.append([label,str(row['step']),fmt(row['epoch'],3),fmt(row['eval_loss'],6),pct(row.get('eval_mean_token_accuracy'),4)])
    table(['Run','Step','Epoch','Eval loss','Eval token accuracy'],rows,[1.7,.7,1,1.7,1.9])
    p('Loss validation giảm ở ba mốc của cả hai stage. Token accuracy là đánh giá dự đoán token khi có prefix mục tiêu, không tương đương chất lượng sinh tự do. Train loss tổng của run có resume và eval loss toàn tập khác phạm vi tích lũy; không dùng riêng tỷ số hai loss để khẳng định overfitting.')
    h('3 2 Chuẩn bị nhãn Pointer và độ dài',2)
    conv=data['pointer_conversion']; lengths=c2['length_report']
    table(['Chỉ tiêu chuyển đổi','Kết quả'],[['CV round-trip đúng CVSchema',f'{conv["round_trip"]["exact_cvschema_matches"]}/{conv["round_trip"]["checked_records"]}'],['Slot trường dài',str(conv['long_field_count'])],['Span contiguous / null',f'{conv["alignment_match_types"]["exact_contiguous"]} / {conv["alignment_match_types"]["absent"]}'],['Alignment mơ hồ',str(conv['ambiguous_alignment_count'])],['Giảm ký tự assistant target',pct(conv['character_impact']['assistant_output_reduction_rate'])],['Tăng ký tự source do đánh số dòng',pct(conv['character_impact']['source_indexing_increase_rate'])]], [4.6,2.4])
    table(['Split','p50','p95','p99','Max','Vượt 5.120'],[[s,*[str(lengths[s][k]) for k in ('p50','p95','p99','max','over_limit')]] for s in ('train','validation')],[1.6,1,1,1,1,1.4])
    p('Độ dài trên tính bằng tokenizer Qwen cho hội thoại Pointer. Giảm ký tự không phải giảm token; lợi ích token khi inference được đo riêng bên dưới. Round-trip hoàn hảo chỉ xác nhận chuyển đổi nhãn nguồn đúng, không xác nhận nhãn nguồn đầy đủ theo quy tắc gold thật.')
    h('4 Quy tắc đánh giá và ý nghĩa chỉ số')
    for text in [
        'Strict JSON: raw output parse trực tiếp. Recoverable JSON: chỉ bỏ wrapper đã cho phép trong parser đóng băng; không sửa JSON thiếu hoặc bị cắt.',
        'CVSchema-valid: đủ khóa và đúng kiểu. Nó chưa xác nhận dữ kiện đúng PDF. Chính sách Communication-only của bộ tổng hợp được báo riêng; không dùng để loại CV thật schema-valid khỏi tính F1.',
        'Field Micro F1: tập giá trị theo 19 trường, chuẩn hóa Unicode NFKC, casefold và khoảng trắng. Không tự dịch, gộp synonym hoặc bỏ dấu câu. Tên, email và điện thoại đã mask không thuộc 19 trường đo.',
        'Precision = TP/(TP+FP), recall = TP/(TP+FN), F1 = 2TP/(2TP+FP+FN). Output schema-invalid được tính như dự đoán rỗng; không bỏ CV đó khỏi mẫu.',
        'Macro F1 chỉ trên field có gold support. Trường không có ví dụ dương có thể không có F1 áp dụng; không được xem nhãn luôn rỗng là bằng chứng trích xuất tốt.',
        'Document exact: toàn cấu trúc nội dung khớp theo cách chuẩn hóa của evaluator tương ứng. F1 52,81% không có nghĩa 52,81% số CV hoàn toàn đúng.',
        'Span contract kiểm tra kiểu và giới hạn dòng. Span exact so với gold đo chọn dòng đúng. Dựng đúng văn bản của một span được chọn không có nghĩa span đó bao phủ thông tin đúng.']:
        bullet(text)
    p('Parser Full JSON v1 và Pointer v1 đã khóa trước test tổng hợp. Một số quy tắc wrapper được phát triển trên validation, vì vậy validation là tập phát triển; GOLD40 được đánh giá hậu kiểm và đã dùng để nghiên cứu lỗi. Tập xác nhận cho cải tiến tiếp theo phải có dữ liệu mới.')
    h('5 Kết quả validation tổng hợp')
    def synthetic_table(split):
        s=data['synthetic'][split]; b=s['fulljson']['models']['base']; a=s['fulljson']['models']['adapter']; q=s['pointer']; n=q['sample_count']
        def full_long(r):
            fs=r['field_metrics']['by_field']; total=[sum(fs[f][k] for f in ('summary','experience.description','projects.details')) for k in ('tp','fp','fn')];return 2*total[0]/(2*total[0]+total[1]+total[2])
        table(['Chỉ tiêu','B0 Full JSON','B1 Full JSON','P0 Pointer'],[
            ['Strict JSON',frac(b['strict_json_count'],n),frac(a['strict_json_count'],n),frac(q['counts']['strict_json'],n)],
            ['Recoverable JSON',frac(b['recoverable_json_count'],n),frac(a['recoverable_json_count'],n),frac(q['counts']['recoverable_json'],n)],
            ['CVSchema-valid',frac(b['schema_valid_count'],n),frac(a['schema_valid_count'],n),frac(q['counts']['cvschema_valid'],n)],
            ['Policy tổng hợp hợp lệ',frac(b['skill_policy_valid_count'],n),frac(a['skill_policy_valid_count'],n),frac(q['counts']['skill_policy_valid'],n)],
            ['Document exact',frac(b['exact_full_json_count'],n),frac(a['exact_full_json_count'],n),frac(q['counts']['cvschema_exact'],n)],
            ['Field Precision',pct(b['field_metrics']['micro']['precision']),pct(a['field_metrics']['micro']['precision']),pct(q['field_metrics']['micro']['precision'])],
            ['Field Recall',pct(b['field_metrics']['micro']['recall']),pct(a['field_metrics']['micro']['recall']),pct(q['field_metrics']['micro']['recall'])],
            ['Field Micro F1',pct(b['field_metrics']['micro']['f1']),pct(a['field_metrics']['micro']['f1']),pct(q['field_metrics']['micro']['f1'])],
            ['Macro F1 field có support',pct(b['field_metrics']['macro_f1_supported_fields']),pct(a['field_metrics']['macro_f1_supported_fields']),pct(q['field_metrics']['macro_f1_supported_fields'])],
            ['Long Text Micro F1',pct(full_long(b)),pct(full_long(a)),pct(q['long_text_micro']['f1'])],
            ['Chạm giới hạn token',str(b['hit_max_new_tokens_count']),str(a['hit_max_new_tokens_count']),str(q['generation']['hit_max_new_tokens_count'])]], [2.65,1.45,1.45,1.45])
        rows=[]
        for name,e in [('B0',s['fulljson_efficiency']['B0']),('B1',s['fulljson_efficiency']['B1']),('P0',q['efficiency'])]:
            rows.append([name,str(e['generated_tokens_total']),fmt(e['generated_tokens_mean']),fmt(e['runtime_seconds_mean']),fmt(e['runtime_seconds_p50']),fmt(e['runtime_seconds_p95'])])
        table(['Model','Output token tổng','Token/CV','Giây/CV','p50 giây','p95 giây'],rows,[.6,1.5,1.2,1.3,1.2,1.2])
        sp=q['span_metrics']['overall'];p(f'Pointer span exact cho slot gold khác null: {sp["non_null_exact"]}/{sp["non_null_gold"]}. Chỉ số này đọc theo phạm vi span của evaluator, tách khỏi schema gate và Document exact.')
    synthetic_table('validation')
    old=data['historical_validation_before_frozen_parser'];p(f'Kết quả B1 ban đầu ngày 27/09 dùng checker cũ: recoverable {old["recoverable_json_count"]}/75, document exact {old["exact_full_json_count"]}/75, Micro F1 {pct(old["field_metrics"]["__micro_all_fields__"]["f1"])}. Bốn wrapper </tool_call> được checker cũ từ chối. Bảng chính dùng parser v1 đã thống nhất, đạt recoverable 75/75 và F1 99,60%. Hai con số thể hiện thay đổi parser trên cùng validation, không phải hai lần train hoặc cải thiện model mới.')
    p('Trên validation Pointer, validator không có retry_targets nên cổng P1 không gọi thêm model và cho kết quả hiệu dụng bằng P0. Không có một run P1 độc lập để chứng minh cải thiện trên tập này. So với B1, P0 tăng Document exact từ 61/75 lên 64/75.')
    h('6 Kết quả test tổng hợp đã khóa')
    synthetic_table('test')
    p('Trên test tổng hợp, B1 tăng Micro F1 từ 87,98% lên 99,52% so với B0. P0 đạt 99,19%, thấp hơn B1 0,34 điểm phần trăm; Document exact tăng từ 57/75 lên 62/75. Mẫu rpv_r00826 thiếu activities và awards nên Pointer parser từ chối toàn output. Chưa chạy P1 riêng trên test này.')
    p('B0/B1 dùng giới hạn 1.280 token; P0 test dùng 768 token theo config đã lưu. P0 không chạm giới hạn và tối đa đã sinh 569 token; B0 có 2 mẫu chạm giới hạn. Chênh lệch ngân sách là một giới hạn đối chứng cần công bố, không gọi mọi điều kiện decoding hoàn toàn giống nhau.')
    p('P0 test giảm 38,49% output token trung bình so với B1. Latency trung bình đã ghi giảm khoảng 37,37%; các run ở phiên khác nhau nên đây là mô tả số đo, chưa là thử nghiệm kiểm soát để xác nhận tốc độ do kiến trúc.')
    h('7 Kết quả trên 80 CV thật')
    r80=data['real80_p1'];eff=r80['efficiency']
    table(['Chỉ tiêu','P0','P1 hiệu dụng'],[
        ['CVSchema-valid',frac(r80['p0_counts']['cvschema_valid'],80),frac(r80['p1_effective_counts']['cvschema_valid'],80)],
        ['Model calls',str(eff['p0_model_calls']),str(eff['total_model_calls'])],
        ['Output token tổng',str(eff['p0_generated_tokens']),str(eff['effective_generated_tokens'])],
        ['Thời gian sinh phút',fmt(eff['p0_runtime_seconds']/60),fmt(eff['effective_runtime_seconds']/60)],
        ['CV còn bị từ chối','9/80','2/80']], [3,2,2])
    p('Nguồn đầu vào: 79 CV dùng text PDF, chỉ EN047 dùng PaddleOCR tiếng Anh. Tất cả 20 VI dùng embedded text. P1 retry 9 CV, khôi phục 7/9, còn EN011 và VI006 cần review. P0 valid EN 56/60, VI 15/20; P1 valid EN 59/60, VI 19/20.')
    p('P1 tăng 8,75 điểm phần trăm schema validity, phát sinh thêm 11,25% model calls, 12,93% output token và 13,29% runtime sinh. Cả 9 retry kết thúc EOS, không chạm 1.280 token. Bộ 80 chưa có nhãn nội dung cho toàn bộ nên không báo F1/accuracy nội dung trên đủ 80; GOLD40 là nhóm có gold bên dưới.')
    h('8 Kết quả GOLD40 thật')
    g=data['gold40']['models']
    table(['Chỉ tiêu','B0','B1','P0','P1'],[
        ['CVSchema-valid',*[frac(g[m]['all']['structural_counts']['cvschema_valid'],40) for m in g]],
        ['Precision',*[pct(g[m]['all']['micro']['precision']) for m in g]],
        ['Recall',*[pct(g[m]['all']['micro']['recall']) for m in g]],
        ['Micro F1',*[pct(g[m]['all']['micro']['f1']) for m in g]],
        ['Micro F1 EN',*[pct(g[m]['en']['micro']['f1']) for m in g]],
        ['Micro F1 VI',*[pct(g[m]['vi']['micro']['f1']) for m in g]],
        ['Macro F1 supported',*[pct(g[m]['all']['macro_f1_supported_fields']) for m in g]],
        ['Long Text F1',*[pct(g[m]['all']['groups']['long_text']['f1']) for m in g]],
        ['Short Field F1',*[pct(g[m]['all']['groups']['short_fields']['f1']) for m in g]],
        ['Document exact',*[f'{g[m]["all"]["document_exact_count"]}/40' for m in g]],
        ['TP / FP / FN',*[f'{g[m]["all"]["micro"]["tp"]}/{g[m]["all"]["micro"]["fp"]}/{g[m]["all"]["micro"]["fn"]}' for m in g]]],[2.4,1.15,1.15,1.15,1.15])
    p('B0 có Field Micro F1 cao nhất trên GOLD40. P1 khôi phục 5 output bị từ chối P0 và tăng recall; precision giảm. P1 chưa cải thiện mọi tiêu chí: Long Text F1 giảm từ 17,39% xuống 16,60%, Document exact vẫn 0. B0 không đứng đầu tất cả tiêu chí entity; bảng entity riêng trong phụ lục tránh suy rộng kết luận từ Micro F1.')
    table(['Cấu hình','F1 điểm','CI 95 phần trăm'],[[m,pct(g[m]['all']['micro']['f1']),' đến '.join(pct(v) for v in diag['real_f1_cv_bootstrap_ci95']['models'][m])] for m in g],[1.8,1.6,3.6])
    p('Khoảng trên dùng bootstrap 5.000 lần, seed 42, phân tầng 20 EN/20 VI và lấy lại toàn CV. Không xem từng field là mẫu độc lập. Khoảng giữ cố định nhãn và nhóm quan sát; chưa bao gồm sai số nhãn hoặc thiên lệch chọn mẫu.')
    rows=[]
    for pair,row in data['gold40']['paired_comparisons'].items():
        rows.append([pair.replace('_minus_',' so với '),fmt(row['delta_micro_f1']*100),' đến '.join(fmt(v*100) for v in row['delta_micro_f1_ci95'])])
    table(['So sánh ghép cặp','Delta F1 điểm phần trăm','CI 95 điểm phần trăm'],rows,[2,2.4,2.6])
    p('Bootstrap ghép cặp này dùng 2.000 lần, seed 42, cũng phân tầng theo ngôn ngữ. P1 tốt hơn P0 trên nhóm quan sát, nhưng P1 vẫn thấp hơn B0. P0 so với B1 có khoảng chứa 0; chưa có bằng chứng rõ về cải thiện F1. Đây là kết quả hậu kiểm, không suy rộng trực tiếp sang quần thể ứng viên.')
    h('8 1 Chi phí suy luận trên cùng GOLD40',2)
    rows=[]
    for m in g:
        e=g[m]['all']['efficiency'];rows.append([m,str(e['model_calls']),str(e['input_tokens_total']),str(e['generated_tokens_total']),str(e['input_tokens_total']+e['generated_tokens_total']),fmt(e['model_seconds_total']/60)])
    table(['Model','Calls','Input token','Output token','Tổng token','Phút sinh'],rows,[.65,.65,1.45,1.45,1.45,1.35])
    rows=[]
    for m in g:
        e=g[m]['all']['efficiency'];rows.append([m,fmt(e['model_seconds_mean']),fmt(e['model_seconds_p50']),fmt(e['model_seconds_p95']),str(e['ended_with_eos_calls']),str(e['hit_max_new_tokens_calls'])])
    table(['Model','Giây/CV','p50 giây','p95 giây','EOS calls','Hit max calls'],rows,[.65,1.3,1.3,1.3,1.25,1.25])
    p('Tất cả cấu hình GOLD40 dùng greedy và giới hạn 1.280 token. P1 gồm 40 lượt đầu và 6 retry thuộc GOLD40; tổng runtime và token bao gồm các lượt không thành công. P1 giảm 48,28% output token so với B0, nhưng tổng input+output tăng 8,27%; không kết luận tổng chi phí giảm từ riêng số output. Model seconds mean của P1 là thời gian đầu+retry chia 40 CV, không phải trung bình thời gian một call. Chưa quy đổi thành USD vì chạy Colab và chưa có giá hạ tầng đo được.')
    h('9 Kết quả OCR pilot đã đo')
    p('Pilot dùng 5 trang VI001, VI006, VI013, VI018, VI020, raster 200 DPI từ PDF có text. Một CV EN047 chỉ có ảnh được kiểm tra riêng. Các trang VI này không phải bộ scan máy ảnh/máy quét mới; text embedded dùng làm tham chiếu nên phép đo OCR vẫn cần công bố giới hạn tham chiếu.')
    rows=[]
    for m,row in data['ocr_hybrid']['summaries'].items():
        rows.append([m,pct(row['token_f1_macro']),pct(row['accented_token_recall_macro']),pct(row['technical_token_recall_macro']),f'{row["critical_exact_matches"]}/{row["critical_reference_values"]}',fmt(row['seconds_per_page_macro'])])
    table(['OCR recognizer','Token F1','Recall dấu','Recall kỹ thuật','Định danh exact','Giây/trang'],rows,[2.5,.9,.9,.9,.9,.9])
    p('PaddleOCR v5 đo detector+recognizer; VietOCR và hybrid chỉ đo recognizer với polygon detector tái sử dụng. Không so trực tiếp runtime end-to-end giữa các hàng. VietOCR không được fine-tune trong đề tài này; đây là đánh giá pretrained recognizer. Các định danh có 15 giá trị tham chiếu, không phải 15 CV.')
    table(['OCR recognizer','CER có thứ tự','WER có thứ tự'],[[m,pct(row['cer_order_sensitive_macro']),pct(row['wer_order_sensitive_macro'])] for m,row in data['ocr_hybrid']['summaries'].items()],[3.6,1.7,1.7])
    p('Đối chứng PaddleOCR v3 Latin đạt Token F1 32,62%, recall token có dấu 7,68%, CER 15,67%, WER 69,78%, 24,15 giây/trang. V5 Latin đạt CER 9,15%, WER 44,56% và confidence macro 96,08%; confidence cao không bảo đảm dấu tiếng Việt đúng. Dictionary audit ghi thiếu 45/74 ký tự VI kiểm tra ở v5 Latin và 53/74 ở v3 Latin.')
    p('Nhánh pilot được chọn cho scan VI là Paddle detector + VietOCR Seq2Seq + Paddle contact fallback. Scan EN dùng PaddleOCR English có cờ review; EN047 chưa có ground truth OCR đầy đủ để báo CER/F1 EN tương đương bảng VI. Benchmark trích xuất real80/GOLD40 không đánh giá nhánh OCR VI vì các PDF VI đều đi tuyến native text.')
    h('10 Phân tích lỗi và giới hạn của dữ liệu')
    train=diag['dataset_coverage']['train']; gold_counts=diag['gold_field_nonempty_cv_counts']
    table(['Độ phủ nhãn','Train 850','GOLD40'],[['CV tiếng Việt','0','20'],['GitHub có giá trị','0',str(gold_counts['personal_info.github_url'])],['GPA có giá trị','0',str(gold_counts['education.gpa'])],['Activities có giá trị','0',str(gold_counts['activities'])],['Awards có giá trị','0',str(gold_counts['awards'])],['Soft skill target','Chỉ Communication','Nhiều kỹ năng tường minh']], [3,2,2])
    p('Lệch phạm vi nhãn đã xác nhận: converter tổng hợp chỉ đưa source.skills sang hard/soft skills; công nghệ dự án giữ ở projects.technologies, Languages không chuyển vào skills. Gold yêu cầu kỹ năng tường minh toàn CV. 543/544 CV train có công nghệ dự án thiếu ít nhất một giá trị nguyên dạng trong hard_skills target, và 850/850 nguồn có Languages. Cần thống nhất phạm vi nhãn trước run mới.')
    p('B1 hard-skills F1 13,69%, B0 58,33%. Heuristic tìm 304 cặp ứng viên trên 26 CV có kỹ năng gold bị chấm thiếu nhưng nằm trong chuỗi kỹ năng B1 dài hơn. Đây là ca nghi ngờ gộp, có thể có ứng viên giả như C/C#/C++; không cộng lại thành TP. Nếu loại riêng hard_skills khỏi phép tổng hợp chẩn đoán, B0 F1 69,28%, B1 62,48%. Điểm chính vẫn giữ đủ 19 trường.')
    p('B1 schema-invalid gồm EN012 lỗi JSON, EN035 thiếu skills.soft_skills và VI006/VI008/VI010 activities sai kiểu. Trên 35 CV cả B0/B1 đều hợp lệ, B0 F1 65,91%, B1 52,28%; lỗi parse không giải thích hết khoảng cách. Phân tích có điều kiện này có thiên lệch chọn mẫu và không thay metric toàn 40.')
    p('GOLD40 nhỏ làm ước lượng và độ phủ hạn chế, không trực tiếp gây overfitting. Các loss validation giảm và test tổng hợp cao chưa cho thấy train quá lâu theo bằng chứng đã lưu; vẫn có khả năng chuyên biệt vào miền/quy tắc tổng hợp. Nhãn gold có thể có lỗi nội dung, nhưng chưa có kết quả rà đủ bộ để lượng hóa ảnh hưởng. Schema validation không thay thế đối chiếu PDF.')
    h('11 Kết luận để trình bày với giảng viên')
    p('Nhóm đã thực hiện huấn luyện QLoRA hai giai đoạn trên Qwen3 4B và có checkpoint, adapter, log và phép đánh giá có truy vết. Stage 1 cải thiện Field Micro F1 trên test tổng hợp từ 87,98% lên 99,52%. Stage 2 Pointer giảm lượng output sinh và tăng Document exact trên test tổng hợp, nhưng Micro F1 giảm nhẹ xuống 99,19%.')
    p('Trên GOLD40 thật, Qwen gốc đạt 65,52% F1, cao hơn adapter Full JSON 48,28%, Pointer một lượt 49,66% và Pointer retry 52,81%. P1 cải thiện tính hợp lệ và recall so với P0, nhưng chưa vượt Qwen gốc về chất lượng nội dung. Do đó chưa có căn cứ tuyên bố fine-tuning hoặc Pointer cải thiện F1 CV thật so với baseline gốc.')
    p('Ưu tiên tiếp theo là rà gold và nguồn PDF, thống nhất phạm vi nhãn train–gold–prompt, bổ sung ví dụ VI/các trường thiếu và đánh giá cấu hình mới trên dữ liệu phát triển riêng. Sau khi chốt cải tiến, cần CV mới để xác nhận. Nếu sửa gold, chấm lại cùng raw output và báo cả gold v1/v2 với lý do sửa; không đưa GOLD40 vào train rồi báo chính nó là test.')
    p('Phần kiến trúc Pointer kế thừa nghiên cứu trước. Các điều chỉnh của nhóm gồm Hybrid Pointer một lượt, hợp đồng xác thực và retry chọn CV lỗi trong giới hạn tài nguyên. Hiệu quả của từng điều chỉnh phải gắn với phép đo cụ thể; không dùng kết quả bài báo khác như kết quả của đề tài.')
    h('12 Các số liệu còn cần bổ sung')
    for text in ['Kết quả rà nhãn nội dung đầy đủ và mức đồng thuận mới của hai người rà; không thể truy dựng IAA lịch sử.', 'Gold cho 40 CV thật còn lại hoặc nhóm CV mới để đo nội dung; real80 hiện chủ yếu có structural benchmark.', 'OCR trên scan VI thực tế, OCR EN có ground truth và kiểm tra reading order độc lập.', 'So checkpoint sớm/cuối trên dev thật mới nếu muốn kiểm chứng over-training; phép đo hiện chưa xác định nguyên nhân nhân quả.', 'Latency và peak VRAM inference trong cùng phiên/hạ tầng, gồm OCR, tải model, cold start và xử lý PDF nếu muốn kết luận end-to-end.', 'Benchmark riêng cho matching, chấm điểm và gợi ý do các thành viên phụ trách.']:
        bullet(text)
    h('Phụ lục A F1 từng trường trên validation tổng hợp')
    def field_table(split):
        s=data['synthetic'][split];b=s['fulljson']['models']['base']['field_metrics']['by_field'];a=s['fulljson']['models']['adapter']['field_metrics']['by_field'];q=s['pointer']['field_metrics']['by_field']
        table(['Trường','Gold support','B0 F1','B1 F1','P0 F1'],[[f,str(b[f]['support']),pct(b[f]['f1']),pct(a[f]['f1']),pct(q[f]['f1'])] for f in b],[2.65,1,.95,.95,1.45])
    field_table('validation')
    h('Phụ lục B F1 từng trường trên test tổng hợp')
    field_table('test')
    h('Phụ lục C F1 từng trường và entity trên GOLD40')
    fs=g['B0']['all']['by_field']
    table(['Trường','Support','B0 F1','B1 F1','P0 F1','P1 F1'],[[f,str(fs[f]['support']),*[pct(g[m]['all']['by_field'][f]['f1']) for m in g]] for f in fs],[2.4,.6,1,1,1,1])
    table(['Entity exact F1','B0','B1','P0','P1'],[[s,*[pct(g[m]['all']['entity_exact'][s]['f1']) for m in g]] for s in ('experience','projects','education')],[2.4,1.15,1.15,1.15,1.15])
    p('Field set F1 tổng hợp các giá trị theo trường, còn entity exact yêu cầu thuộc cùng object và khớp đủ thuộc tính; hai phép đo trả lời khác nhau. Toàn bộ TP/FP/FN, precision/recall/F1 từng trường, breakdown EN/VI và log train nằm trong tệp số liệu JSON kèm báo cáo.')
    h('Phụ lục D Nguồn số liệu và tài liệu nghiên cứu')
    for path in ['training/final_adapter-full_1.zip','training/qwen3_4b_cvschema2_pointer_stage2_full_v1_final_adapter.zip','training/evaluation/results/validation_base_vs_adapter_v1/comparison_summary.json','training/evaluation/results/locked_test_base_vs_adapter_v1/comparison_summary.json','training/evaluation/results/pointer_stage2_validation_p0_v1/evaluation_summary.json','training/evaluation/results/pointer_stage2_locked_test_p0_v1/evaluation_summary.json','training/evaluation/results/pointer_real80_selective_retry_p1_v1/evaluation_summary.json','training/evaluation/results/gold40_b0_b1_p0_p1_v1/summary.json','training/audits/gold40_generalization_2026-10-11_v2/diagnostics.json','experiments/ocr/vietocr_hybrid_pilot_v1/comparison_summary.json']:
        bullet(path)
    p('Tệp số liệu kèm báo cáo lưu hash đầy đủ của nguồn, contract train, log đã ghi và metric gốc. Kiểm tra lại CRC và 10/10 manifest member của mỗi adapter stage 1/smoke/full; kiểm tra manifest kết quả GOLD40, dataset và hash gold. Báo cáo không thay raw output, adapter, parser, nhãn hoặc kết quả cũ.')
    p('Hu và cộng sự, LoRA, 2021: https://arxiv.org/abs/2106.09685. LoRA giữ cố định trọng số base và học các ma trận thích nghi hạng thấp; đây là căn cứ chọn adapter thay vì cập nhật toàn base.')
    p('Dettmers và cộng sự, QLoRA, 2023: https://arxiv.org/abs/2305.14314. QLoRA kết hợp base lượng tử hóa 4-bit với LoRA, NF4 và double quantization để giảm yêu cầu bộ nhớ. VRAM đo trong đề tài xác nhận cấu hình hiện tại chạy được, không chứng minh chất lượng CV thật.')
    p('Layout-Aware Parsing Meets Efficient LLMs, 2025: https://arxiv.org/abs/2510.09722. Pointer là hướng kế thừa từ bài báo đã đọc; khác dữ liệu và giao thức đánh giá nên không so trực tiếp F1 của nhóm với số báo cáo trong bài.')
    return blocks

def markdown(blocks):
    result=[]
    for block in blocks:
        if block[0]=='heading':
            result.append('#'*max(1,block[1]+1)+' '+block[2])
        elif block[0]=='paragraph': result.append(block[1])
        elif block[0]=='bullet': result.append('- '+block[1])
        else:
            headers,rows=block[1:3]
            clean=lambda x:str(x).replace('|','/').replace('\n','<br>')
            result.append('| '+' | '.join(map(clean,headers))+' |\n| '+' | '.join('---' for _ in headers)+' |\n'+'\n'.join('| '+' | '.join(map(clean,row))+' |' for row in rows))
    return '\n\n'.join(result)+'\n'

def word(blocks,path):
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.72)
    section.left_margin=section.right_margin=Inches(.75)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet'):
        s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0);s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.14;normal.paragraph_format.space_after=Pt(7)
    doc.styles['Title'].font.size=Pt(23)
    doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
    doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
    doc.core_properties.title='Dữ liệu huấn luyện và đánh giá mô đun trích xuất CV'
    doc.core_properties.author='Nhóm CNTT KLCN029'
    for block in blocks:
        if block[0]=='heading':
            if block[1]==0:doc.add_paragraph(block[2],'Title')
            else:doc.add_heading(block[2],block[1])
        elif block[0] in ('paragraph','bullet'):
            para=doc.add_paragraph(block[1],style='List Bullet' if block[0]=='bullet' else 'Normal')
            if block[1].startswith(('training/','experiments/')):
                for run in para.runs:run.font.size=Pt(9)
        else:
            headers,rows,widths=block[1:]
            t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            if widths:
                for column,width in zip(t.columns,widths):column.width=Inches(width)
            for j,text in enumerate(headers):t.rows[0].cells[j].text=str(text)
            for row in rows:
                cells=t.add_row().cells
                for j,text in enumerate(row):cells[j].text=str(text)
            properties=t._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            properties.append(borders)
            for i,row in enumerate(t.rows):
                no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                if i==0:
                    header=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(header)
                for j,cell in enumerate(row.cells):
                    if widths:cell.width=Inches(widths[j])
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr();fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'17365D' if i==0 else ('EEF3F8' if i%2==0 else 'FFFFFF'));pr.append(fill)
                    margins=OxmlElement('w:tcMar')
                    for name,value in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
                        e=OxmlElement('w:'+name);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
                    pr.append(margins)
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.05
                        para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                        for run in para.runs:run.font.size=Pt(9.5);run.font.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
            spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(3);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1
            spacer.add_run().font.size=Pt(3)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('CNTT KLCN029 | Trang ').font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(path)

if __name__=='__main__':
    data=collect();blocks=report(data)
    (OUT/(STEM+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (OUT/(STEM+'.md')).write_text(markdown(blocks),encoding='utf-8')
    word(blocks,OUT/(STEM+'.docx'))
    (OUT/(STEM+'_MANIFEST.json')).write_text(json.dumps({'date':'2026-10-11','docx_visual_qa':'NOT_VERIFIED; packaged renderer cannot locate LibreOffice soffice.exe','sources':SOURCES,'outputs':{suffix:sha((OUT/(STEM+suffix)).read_bytes()) for suffix in ('.md','.docx','.json')}},indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'report':str(OUT/(STEM+'.docx')),'sources':len(SOURCES),'blocks':len(blocks)},ensure_ascii=True))
