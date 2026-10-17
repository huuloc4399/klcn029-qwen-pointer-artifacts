from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
ASSET_DIR = OUT_DIR / "assets"
OUTPUT = OUT_DIR / "Bao_cao_tien_do_OCR_va_trich_xuat_KLCN029.docx"

BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
MID_GRAY = "D9D9D9"
WHITE = "FFFFFF"


def font_path(bold: bool = False) -> str:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    raise FileNotFoundError("Không tìm thấy font Windows để tạo sơ đồ")


def create_architecture_diagram(path: Path) -> None:
    width, height = 1800, 930
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(font_path(True), 46)
    box_font = ImageFont.truetype(font_path(False), 30)
    small_font = ImageFont.truetype(font_path(False), 25)
    draw.text((width // 2, 40), "Luồng xử lý đề xuất cho baseline toàn bộ", fill="#000000", font=title_font, anchor="ma")

    def box(x1: int, y1: int, x2: int, y2: int, text: str, fill: str = "#EAF2F8", bold: bool = False) -> None:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=fill, outline="#17365D", width=4)
        f = ImageFont.truetype(font_path(bold), 29 if bold else 27)
        lines = text.split("\n")
        line_height = 38
        start_y = (y1 + y2 - line_height * len(lines)) / 2 + 4
        for index, line in enumerate(lines):
            draw.text(((x1 + x2) / 2, start_y + index * line_height), line, fill="#000000", font=f, anchor="ma")

    def arrow(x1: int, y1: int, x2: int, y2: int) -> None:
        draw.line((x1, y1, x2, y2), fill="#4D4D4D", width=5)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 17
        p1 = (x2 - size * math.cos(angle - 0.55), y2 - size * math.sin(angle - 0.55))
        p2 = (x2 - size * math.cos(angle + 0.55), y2 - size * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), p1, p2], fill="#4D4D4D")

    box(70, 165, 330, 285, "Tải CV\nPDF hoặc ảnh", "#F4F6F7", True)
    box(430, 165, 780, 285, "Kiểm tra lớp text\nvà mật độ pixel", "#EAF2F8", True)
    arrow(330, 225, 430, 225)

    box(890, 95, 1230, 215, "Có text đạt yêu cầu\nPyMuPDF", "#E8F5E9")
    box(890, 285, 1230, 405, "Trang ảnh\nChọn nhánh OCR", "#FFF4E5")
    arrow(780, 205, 890, 155)
    arrow(780, 245, 890, 345)
    draw.text((815, 158), "Có", fill="#000000", font=small_font, anchor="mm")
    draw.text((815, 318), "Không", fill="#000000", font=small_font, anchor="mm")

    box(1360, 250, 1720, 365, "Tiếng Anh\nPaddleOCR English", "#F4F6F7")
    box(1360, 430, 1720, 580, "Tiếng Việt\nPaddle detector\nVietOCR Seq2Seq", "#F4F6F7")
    arrow(1230, 330, 1360, 307)
    arrow(1230, 370, 1360, 505)

    box(900, 650, 1220, 765, "Chuẩn hóa text\nvà reading order", "#EAF2F8")
    box(1330, 650, 1570, 765, "Qwen3 4B\nQLoRA adapter", "#EAF2F8", True)
    box(1620, 650, 1780, 765, "Parser v1", "#EAF2F8", True)
    arrow(1060, 405, 1060, 650)
    arrow(1540, 580, 1160, 650)
    arrow(1220, 707, 1330, 707)
    arrow(1570, 707, 1620, 707)
    box(1370, 820, 1740, 905, "JSON CVSchema 2.0", "#E8F5E9", True)
    arrow(1700, 765, 1625, 820)
    image.save(path, dpi=(200, 200))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="8") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_column_width(cell, inches: float) -> None:
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_table(table, widths: list[float], numeric_columns: set[int] | None = None) -> None:
    numeric_columns = numeric_columns or set()
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, cell in enumerate(row.cells):
            set_column_width(cell, widths[col_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, NAVY if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else WHITE))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if row_index == 0 or col_index in numeric_columns:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else BLACK)
                    if row_index == 0:
                        run.bold = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], numeric_columns: set[int] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    format_table(table, widths, numeric_columns)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_run_font(run, name="Times New Roman", size=11, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0.28)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.2
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc: Document, text: str, level: int = 0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_number(doc: Document, text: str, number: int):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run_font(paragraph.add_run(f"{number}. "), bold=True)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_caption(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(9)
    set_run_font(paragraph.add_run(text), size=9.5, italic=True)
    return paragraph


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)

    for section in doc.sections:
        section.header.is_linked_to_previous = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(header.add_run("KLCN029  Báo cáo tiến độ OCR và trích xuất CV"), size=8.5, color="555555")
        add_page_field(section.footer.paragraphs[0])


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_source(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    set_run_font(paragraph.add_run("Nguồn số liệu: "), size=9, bold=True, color="555555")
    set_run_font(paragraph.add_run(text), size=9, italic=True, color="555555")


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    architecture_path = ASSET_DIR / "ocr_extraction_architecture.png"
    create_architecture_diagram(architecture_path)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Báo cáo tiến độ OCR và mô hình trích xuất CV KLCN029"
    doc.core_properties.subject = "Báo cáo gửi giảng viên hướng dẫn trước khi chạy baseline toàn bộ"
    doc.core_properties.author = "Nhóm sinh viên KLCN029"

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP HỒ CHÍ MINH"), size=12, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("NGÀNH CÔNG NGHỆ THÔNG TIN"), size=12, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    direct_ppr = title._p.get_or_add_pPr()
    direct_border = direct_ppr.find(qn("w:pBdr"))
    if direct_border is not None:
        direct_ppr.remove(direct_border)
    title.add_run("BÁO CÁO TIẾN ĐỘ OCR VÀ MÔ HÌNH TRÍCH XUẤT CV")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    set_run_font(p.add_run("Đề tài KLCN029"), size=14, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    set_run_font(p.add_run("Nghiên cứu hệ thống hỗ trợ đánh giá và cải thiện CV sử dụng mô hình ngôn ngữ lớn"), size=13, italic=True)

    cover_rows = [
        ["Giảng viên hướng dẫn", "Huỳnh Thị Cẩm Dung"],
        ["Sinh viên thực hiện", "Nguyễn Hữu Lộc  2001230468  14DHTH01\nDương Chí Khải  2001230364  14DHTH01\nLê Tâm Đắc  2001230169  14DHTH03"],
        ["Thời điểm báo cáo", "Ngày 29 tháng 09 năm 2026"],
        ["Mục đích", "Báo cáo kết quả pilot và xin ý kiến trước khi chạy baseline end to end trên 80 CV thật"],
    ]
    table = add_table(doc, ["Thông tin", "Nội dung"], cover_rows, [1.85, 5.0])
    table.rows[0]._element.getparent().remove(table.rows[0]._element)
    for row_idx, row in enumerate(table.rows):
        shade_cell(row.cells[0], PALE_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("TP Hồ Chí Minh  2026"), size=11)
    doc.add_page_break()

    add_heading(doc, "Tóm tắt kết quả và nội dung xin hướng dẫn", 1)
    add_body(doc, "Nhóm đã hoàn thành ba phần nền tảng của nhánh Resume Parsing: thiết kế CVSchema 2.0 và chính sách kỹ năng; huấn luyện Qwen3 4B bằng QLoRA; xây dựng Parser v1 và đánh giá trên validation cùng test khóa. Đồng thời, nhóm đã kiểm tra nhánh OCR trên CV thật để lựa chọn engine cho tiếng Anh và tiếng Việt.")
    add_body(doc, "Kết quả test khóa trên 75 CV tổng hợp tiếng Anh cho thấy Qwen sau QLoRA đạt Micro F1 99,52%, so với 87,98% của Qwen gốc. Toàn bộ 75 đầu ra của adapter được Parser v1 khôi phục và vượt CVSchema 2.0. Tuy nhiên, raw output của adapter có wrapper suy luận nên strict JSON bằng 0%; Parser v1 là thành phần bắt buộc trong pipeline.")
    add_body(doc, "Pilot OCR tiếng Việt cho thấy PaddleOCR Latin không giữ được dấu tốt. Phương án lai gồm PaddleOCR phát hiện vùng chữ, VietOCR Seq2Seq nhận dạng nội dung tiếng Việt và PaddleOCR nhận dạng lại dòng liên hệ đạt Token F1 94,66%, recall từ có dấu 96,57% và 13/15 định danh đúng tuyệt đối.")
    add_body(doc, "Nhóm đề nghị giảng viên xem xét kiến trúc OCR theo ngôn ngữ, phạm vi dữ liệu cần bổ sung và kế hoạch đánh giá end to end trước khi nhóm chạy baseline trên toàn bộ 80 CV thật.")

    add_heading(doc, "Kết luận hiện tại", 2)
    add_bullet(doc, "PDF có lớp văn bản: dùng PyMuPDF làm nhánh chính.")
    add_bullet(doc, "CV scan tiếng Anh: dùng PaddleOCR English và gắn cờ review cho lỗi ký tự hoặc thứ tự hai cột.")
    add_bullet(doc, "CV scan tiếng Việt: dùng PaddleOCR detector, VietOCR vgg_seq2seq và fallback PaddleOCR cho email, URL, số điện thoại.")
    add_bullet(doc, "Trích xuất cấu trúc: dùng Qwen3 4B với adapter QLoRA, sau đó qua Parser v1 và kiểm tra CVSchema 2.0.")
    add_bullet(doc, "Chưa kết luận chất lượng end to end trên 80 CV thật vì chưa chạy toàn bộ pipeline và chưa có gold JSON khóa cho tập đánh giá thực tế.")

    add_heading(doc, "Phạm vi báo cáo", 2)
    add_body(doc, "Báo cáo tập trung vào OCR và trích xuất thông tin CV. Các phần chấm điểm CV, so khớp CV với Job Description, RAG và sinh gợi ý cải thiện chưa nằm trong đợt thực nghiệm này.")

    add_heading(doc, "Kiến trúc pipeline đề xuất", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(architecture_path), width=Inches(6.65))
    add_caption(doc, "Hình 1  Luồng xử lý dự kiến cho baseline end to end")
    add_body(doc, "Nút chọn ngôn ngữ trên giao diện dự kiến gồm Tự động, Tiếng Việt và Tiếng Anh. Lựa chọn ngôn ngữ chỉ điều khiển nhánh OCR khi trang không có lớp text đủ dùng. Với PDF có lớp text, pipeline giữ dữ liệu gốc từ PyMuPDF nhằm tránh lỗi OCR không cần thiết.")

    add_heading(doc, "Dữ liệu và hợp đồng đầu ra", 1)
    add_heading(doc, "Bộ dữ liệu huấn luyện", 2)
    add_body(doc, "Nhóm sử dụng bộ dữ liệu sukhrobnurali resume parsing vision, revision 3c254be39d45f69c6e20bfb325b4e6ccb3b5e393, giấy phép CC BY 4.0. Dữ liệu được chuyển sang CVSchema 2.0 và khóa theo ba tập độc lập.")
    add_table(
        doc,
        ["Tập", "Số mẫu", "Mục đích", "Quy tắc sử dụng"],
        [
            ["Train", "850", "Huấn luyện QLoRA", "Được dùng để cập nhật adapter"],
            ["Validation", "75", "Chọn cấu hình và đóng băng parser", "Không dùng cập nhật trọng số"],
            ["Test khóa", "75", "Đánh giá cuối", "Chỉ mở sau khi đóng băng parser và metric"],
        ],
        [1.05, 0.85, 2.15, 2.85],
        {1},
    )
    add_source(doc, "training/data/resume_parsing_vision_cvschema_v2/manifest.json và conversion_report.json")
    add_body(doc, "Cả validation và test khóa đều là CV tổng hợp tiếng Anh. Ảnh không được tải về; trường resume_text là đầu vào mô hình. Tên, email, số điện thoại và ngày sinh tổng hợp đã được che trong văn bản nguồn. Vì vậy kết quả mô hình chỉ chứng minh năng lực ánh xạ text sạch sang schema, chưa chứng minh OCR hoặc khả năng tổng quát sang CV tiếng Việt thực tế.")

    add_heading(doc, "Bộ 80 CV thật của đề tài", 2)
    add_table(
        doc,
        ["Chỉ số", "Kết quả"],
        [
            ["CV tiếng Anh", "60"],
            ["CV tiếng Việt", "20"],
            ["Tổng số trang", "110"],
            ["Nội dung độc nhất", "77"],
            ["Cặp tệp trùng SHA 256", "3"],
            ["Tệp lỗi hoặc có mật khẩu", "0"],
        ],
        [4.95, 1.95],
        {1},
    )
    add_body(doc, "Rà soát lớp text cho thấy toàn bộ 20 CV tiếng Việt hiện có đều đọc được bằng PyMuPDF. Trong tám trang tiếng Anh từng bị đánh dấu ít text, chỉ EN047 có đủ nội dung ảnh để cần OCR; bảy trang còn lại gần trắng hoặc chỉ có phần tử trang trí.")

    add_heading(doc, "CVSchema 2.0 và chính sách kỹ năng", 2)
    add_body(doc, "Đầu ra gồm thông tin cá nhân, tóm tắt, kỹ năng, kinh nghiệm, dự án, học vấn, chứng chỉ, hoạt động và thành tích. Theo quyết định đã được nhóm rà soát ngày 23 tháng 09 năm 2026, chỉ Communication thuộc soft skill; các kỹ năng nguồn còn lại thuộc hard skill. Parser kiểm tra không cho một kỹ năng xuất hiện đồng thời ở hai nhóm.")
    doc.add_page_break()

    add_heading(doc, "Thử nghiệm OCR", 1)
    add_heading(doc, "Mục tiêu và thiết lập", 2)
    add_body(doc, "Pilot OCR dùng năm trang đầu của VI001, VI006, VI013, VI018 và VI020, raster ở 200 DPI. Lớp text PDF được dùng làm tham chiếu. Cùng polygon phát hiện từ PP OCRv5 mobile detector được tái sử dụng để so sánh recognizer, giúp cô lập ảnh hưởng của phần nhận dạng ký tự.")
    add_body(doc, "Các chỉ số gồm Character Error Rate và Word Error Rate có xét thứ tự, Token F1 dạng multiset, recall token có dấu, recall thuật ngữ kỹ thuật và tỷ lệ email, URL, số điện thoại khớp tuyệt đối. Token F1 và recall từ có dấu từ 90% trở lên được dùng làm ngưỡng pilot cho recognizer.")

    add_heading(doc, "Kết quả so sánh", 2)
    add_table(
        doc,
        ["Phương án", "Token F1", "Recall có dấu", "Recall kỹ thuật", "Định danh đúng", "Giây trên trang"],
        [
            ["PaddleOCR v5 Latin", "55,81%", "39,60%", "95,95%", "13/15", "34,84"],
            ["VietOCR Seq2Seq", "94,15%", "96,57%", "93,58%", "7/15", "7,97*"],
            ["VietOCR Transformer", "96,38%", "99,01%", "91,75%", "6/15", "38,21*"],
            ["Hybrid Seq2Seq", "94,66%", "96,57%", "93,58%", "13/15", "7,97*"],
            ["Hybrid Transformer", "96,82%", "99,01%", "91,75%", "13/15", "38,21*"],
        ],
        [1.75, 0.95, 1.05, 1.05, 1.05, 1.05],
        {1, 2, 3, 4, 5},
    )
    add_source(doc, "experiments/ocr/vietocr_hybrid_pilot_v1/comparison_summary.json. Dấu sao là thời gian recognition only; PaddleOCR gồm detector và recognizer nên không so sánh trực tiếp thời gian end to end.")

    add_heading(doc, "Phân tích kết quả", 2)
    add_body(doc, "PaddleOCR Latin có confidence trung bình cao nhưng thiếu nhiều ký tự tiếng Việt dựng sẵn, dẫn đến mất dấu trên diện rộng. VietOCR cải thiện rõ rệt nội dung tiếng Việt. Transformer cao hơn Seq2Seq 2,23 điểm Token F1 nhưng chậm hơn khoảng 4,8 lần trên CPU, vì vậy Seq2Seq phù hợp hơn với baseline có giới hạn tài nguyên.")
    add_body(doc, "VietOCR thường thay ký tự @ bằng O hoặc 0 và có thể nhầm chữ với số trong URL. Bộ chọn lai dùng VietOCR cho nội dung tiếng Việt và kết quả PaddleOCR cho các dòng liên hệ, tăng định danh đúng tuyệt đối từ 7/15 lên 13/15 đối với Seq2Seq. Hai lỗi còn lại nằm ở URL của VI020.")

    add_heading(doc, "Quyết định OCR tạm thời", 2)
    add_number(doc, "Trích text bằng PyMuPDF trước khi xét OCR.", 1)
    add_number(doc, "Chỉ OCR khi trang có dưới 80 ký tự native và tỷ lệ pixel tối từ 0,02 trở lên; cần cài đặt và kiểm thử rule này trước batch toàn bộ.", 2)
    add_number(doc, "Dùng PaddleOCR English cho CV scan tiếng Anh.", 3)
    add_number(doc, "Dùng Paddle detector, VietOCR vgg_seq2seq và contact fallback cho CV scan tiếng Việt.", 4)
    add_number(doc, "Gắn cờ review khi email, URL hoặc số điện thoại không vượt validation định dạng hoặc có confidence thấp.", 5)

    add_heading(doc, "Giới hạn OCR", 2)
    add_bullet(doc, "Năm trang tiếng Việt được raster từ PDF sạch, chưa đại diện ảnh chụp nghiêng, nhiễu hoặc scan chất lượng thấp.")
    add_bullet(doc, "Benchmark recognizer tái sử dụng polygon detector, nên chưa có thời gian end to end cho nhánh hybrid.")
    add_bullet(doc, "Thứ tự đọc CV nhiều cột vẫn cần module layout riêng và kiểm tra trực quan.")

    add_heading(doc, "Huấn luyện mô hình trích xuất", 1)
    add_heading(doc, "Cấu hình QLoRA", 2)
    add_table(
        doc,
        ["Thành phần", "Cấu hình"],
        [
            ["Mô hình nền", "Qwen Qwen3 4B Instruct 2507"],
            ["Phương pháp", "QLoRA với NF4 và double quantization"],
            ["LoRA", "Rank 16  Alpha 32  Dropout 0,05  toàn bộ linear"],
            ["Dữ liệu", "850 train  75 validation"],
            ["Batch hiệu dụng", "8"],
            ["Số epoch và bước", "1 epoch  107 bước"],
            ["Thời gian Trainer", "2.292,47 giây  khoảng 38 phút 12 giây"],
            ["Peak GPU", "8,184 GiB allocated  12,447 GiB reserved"],
            ["Tham số LoRA", "33.030.144 tham số FP32"],
        ],
        [2.05, 4.85],
    )
    add_source(doc, "training/audits/final_adapter-full_1.audit.md")
    add_body(doc, "Adapter ZIP đã qua kiểm tra CRC, manifest, kích thước và SHA 256. Artifact chứa đủ trọng số, tokenizer, run config, metrics và Trainer state. Cấu hình checkpoint trên Google Drive cho phép tiếp tục huấn luyện khi Colab Free bị ngắt.")

    add_heading(doc, "Dấu hiệu hội tụ", 2)
    add_table(
        doc,
        ["Bước", "Eval loss", "Token accuracy"],
        [["40", "0,015352", "0,995554"], ["80", "0,011234", "0,999367"], ["107", "0,010912", "0,999482"]],
        [2.1, 2.4, 2.4],
        {0, 1, 2},
    )
    add_body(doc, "Eval loss giảm ở cả ba mốc và log không có NaN hoặc Inf. Đây là teacher forced metrics, dùng để kiểm tra quá trình tối ưu; chất lượng sinh tự do được đánh giá riêng bằng validation và test khóa.")

    add_heading(doc, "Parser v1", 2)
    add_body(doc, "Parser v1 được đóng băng sau validation. Parser chỉ loại bỏ wrapper đã khai báo gồm khối think hoàn chỉnh, thẻ đóng tool call ở đầu và JSON fence bao toàn bộ phản hồi; không sửa JSON thiếu hoặc output bị cắt. Thiết kế này tách năng lực mô hình khỏi cơ chế phục hồi định dạng.")

    add_heading(doc, "Đánh giá mô hình trích xuất", 1)
    add_heading(doc, "Kết quả validation", 2)
    add_table(
        doc,
        ["Chỉ số", "Qwen gốc", "Qwen QLoRA", "Chênh lệch"],
        [
            ["Recoverable JSON", "97,33%", "100,00%", "+2,67 điểm %"],
            ["Schema hợp lệ", "97,33%", "100,00%", "+2,67 điểm %"],
            ["Chính sách kỹ năng", "46,67%", "98,67%", "+52,00 điểm %"],
            ["Exact full JSON", "0,00%", "81,33%", "+81,33 điểm %"],
            ["Micro F1", "88,06%", "99,60%", "+11,53 điểm %"],
            ["Macro F1", "77,62%", "98,77%", "+21,15 điểm %"],
        ],
        [2.2, 1.45, 1.55, 1.7],
        {1, 2, 3},
    )
    add_source(doc, "training/evaluation/results/validation_base_vs_adapter_v1/comparison_report.md")

    add_heading(doc, "Kết quả test khóa", 2)
    add_table(
        doc,
        ["Chỉ số", "Qwen gốc", "Qwen QLoRA", "Chênh lệch"],
        [
            ["Recoverable JSON", "97,33%", "100,00%", "+2,67 điểm %"],
            ["Schema hợp lệ", "97,33%", "100,00%", "+2,67 điểm %"],
            ["Chính sách kỹ năng", "53,33%", "100,00%", "+46,67 điểm %"],
            ["Exact full JSON", "1,33%", "76,00%", "+74,67 điểm %"],
            ["Micro precision", "88,39%", "99,55%", "+11,16 điểm %"],
            ["Micro recall", "87,58%", "99,49%", "+11,91 điểm %"],
            ["Micro F1", "87,98%", "99,52%", "+11,54 điểm %"],
            ["Macro F1", "78,58%", "98,85%", "+20,27 điểm %"],
            ["Chạm giới hạn token", "2/75", "0/75", "Giảm 2 mẫu"],
        ],
        [2.2, 1.45, 1.55, 1.7],
        {1, 2, 3},
    )
    add_source(doc, "training/evaluation/results/locked_test_base_vs_adapter_v1/report_ready_results_vi.md")
    add_body(doc, "Micro F1 của adapter trên test chỉ thấp hơn validation 0,08 điểm phần trăm. Điều này cho thấy chất lượng theo trường ổn định trên phần dữ liệu chưa dùng để chọn cấu hình. Exact full JSON giảm 5,33 điểm phần trăm vì chỉ một trường sai cũng làm cả mẫu bị tính sai.")

    add_heading(doc, "Lỗi còn lại trên test khóa", 2)
    add_table(
        doc,
        ["Trường", "Mẫu có sai khác", "F1 test", "Nhận xét"],
        [
            ["projects role", "13", "87,58%", "Vai trò thường phải suy ra từ mô tả"],
            ["skills hard skills", "4", "99,59%", "Một số khác biệt chuẩn hóa danh sách"],
            ["experience job title", "1", "99,82%", "Sai khác đơn lẻ"],
            ["certifications", "1", "95,74%", "Support nhỏ nên nhạy với một lỗi"],
        ],
        [1.65, 1.15, 1.0, 3.1],
        {1, 2},
    )

    add_heading(doc, "Vai trò bắt buộc của Parser", 2)
    add_body(doc, "Strict JSON của raw adapter bằng 0% vì 70 mẫu có wrapper think và 5 mẫu có thẻ đóng tool call ở đầu. Parser v1 khôi phục đủ 75 mẫu và toàn bộ đầu ra sau parser hợp lệ theo schema. Pipeline triển khai phải giữ bước parser; kết quả raw không được chuyển thẳng cho thành phần tiêu thụ JSON.")

    add_heading(doc, "Những kết luận đã có và khoảng trống", 1)
    add_table(
        doc,
        ["Nội dung", "Trạng thái", "Bằng chứng hoặc khoảng trống"],
        [
            ["QLoRA cải thiện extraction", "Đã chứng minh", "Test khóa tăng Micro F1 từ 87,98% lên 99,52%"],
            ["Tuân thủ CVSchema 2.0", "Đã chứng minh", "75/75 test hợp lệ sau Parser v1"],
            ["OCR tiếng Việt trên bản raster sạch", "Đã thử nghiệm", "Hybrid Seq2Seq đạt Token F1 94,66%"],
            ["OCR scan tiếng Anh", "Pilot một CV", "EN047 đọc được nội dung nhưng còn lỗi ký tự và reading order"],
            ["OCR ảnh chụp hoặc scan xấu tiếng Việt", "Chưa chứng minh", "Chưa có tập mẫu đại diện"],
            ["Extraction trên CV tiếng Việt thật", "Chưa chứng minh", "Chưa có gold JSON khóa và chưa chạy toàn bộ"],
            ["Pipeline end to end trên 80 CV", "Chưa chạy", "Cần hoàn thiện router và log thực nghiệm"],
        ],
        [2.05, 1.35, 3.5],
    )

    add_heading(doc, "Kế hoạch chạy baseline toàn bộ", 1)
    add_number(doc, "Đóng băng phiên bản code, schema, parser, adapter và cấu hình OCR; lưu hash cho từng artifact.", 1)
    add_number(doc, "Cài đặt router theo lớp text và mật độ pixel; thêm lựa chọn Tự động, Tiếng Việt và Tiếng Anh trên giao diện.", 2)
    add_number(doc, "Chạy bước trích text hoặc OCR trên 80 CV, lưu engine, thời gian, số dòng, confidence và cờ review theo từng trang.", 3)
    add_number(doc, "Đưa text đã chuẩn hóa vào Qwen QLoRA bằng greedy decoding; lưu raw output trước parser.", 4)
    add_number(doc, "Áp dụng Parser v1, kiểm tra CVSchema 2.0 và chính sách kỹ năng; lưu lý do lỗi thay vì sửa ngầm.", 5)
    add_number(doc, "Báo cáo riêng ba nhóm: English native text, Vietnamese native text và OCR fallback.", 6)
    add_number(doc, "Sau khi có gold JSON, tính Precision, Recall, F1 theo trường, Exact full JSON và tỷ lệ schema hợp lệ; phân loại lỗi OCR, layout, extraction và parser.", 7)

    add_heading(doc, "Artifact cần lưu sau lượt chạy", 2)
    add_table(
        doc,
        ["Artifact", "Nội dung"],
        [
            ["run manifest", "Phiên bản model, adapter, parser, schema, OCR và hash dữ liệu"],
            ["page routing log", "Native chars, dark pixel ratio, ngôn ngữ, engine, runtime, cờ review"],
            ["raw predictions", "Đầu ra Qwen trước parser"],
            ["parsed predictions", "JSON sau Parser v1 và kết quả schema validation"],
            ["metrics", "Tổng hợp toàn bộ, theo ngôn ngữ và theo nhánh OCR"],
            ["error analysis", "Danh sách mẫu lỗi và nguyên nhân được rà soát"],
        ],
        [2.0, 4.9],
    )

    add_heading(doc, "Nội dung xin ý kiến giảng viên", 1)
    add_body(doc, "Trước khi chạy baseline toàn bộ, nhóm kính đề nghị giảng viên góp ý và xác nhận các nội dung sau:")
    questions = [
        "Kiến trúc hai nhánh OCR theo ngôn ngữ có phù hợp không, hay nhóm cần thống nhất một OCR duy nhất để dễ so sánh?",
        "Giảng viên có đồng ý dùng chế độ Tự động làm mặc định và cho người dùng chọn Tiếng Việt hoặc Tiếng Anh khi tải CV không?",
        "Phương án PaddleOCR detector, VietOCR Seq2Seq và fallback dòng liên hệ có được xem là baseline chính cho CV scan tiếng Việt không?",
        "Nhóm có cần bổ sung một tập CV tiếng Việt scan hoặc ảnh chụp thật trước khi khóa kết luận OCR không? Nếu có, số lượng tối thiểu nên là bao nhiêu trang?",
        "Có thể dùng 20 CV thật trong test khóa để gán gold JSON không, và có cần hai người gán nhãn độc lập rồi đối chiếu bất đồng không?",
        "Đối với đóng góp huấn luyện, việc fine tune Qwen bằng QLoRA và so sánh với Qwen gốc đã đủ yêu cầu chưa, hay cần thêm một cấu hình huấn luyện hoặc một model đối chứng?",
        "Việc giữ Parser v1 để loại wrapper định dạng có phù hợp với cách đánh giá không, khi raw output của adapter không đạt strict JSON?",
        "Giảng viên có đồng ý báo cáo tách ba tầng OCR, extraction trên text sạch và end to end để tránh quy lỗi sai tầng không?",
        "Trường projects role hiện là lỗi chính còn lại. Nhóm nên giữ trường này, đổi thành nullable khi CV không nêu rõ hay bổ sung hướng dẫn suy luận trong dữ liệu huấn luyện?",
        "Nhóm có thể bắt đầu chạy baseline toàn bộ 80 CV sau khi hoàn thiện router và đóng băng run manifest không?",
    ]
    for index, question in enumerate(questions, 1):
        add_number(doc, question, index)

    add_heading(doc, "Đề xuất của nhóm", 1)
    add_body(doc, "Nhóm đề xuất giữ kiến trúc hiện tại và thực hiện thêm một kiểm tra nhỏ trên CV tiếng Việt scan thực tế trước khi đóng băng OCR. Sau khi được giảng viên thống nhất phạm vi đánh giá, nhóm sẽ chạy baseline trên 80 CV và báo cáo tách biệt lỗi OCR, lỗi bố cục, lỗi trích xuất ngữ nghĩa và lỗi định dạng đầu ra.")
    add_body(doc, "Trong báo cáo khóa luận, kết quả 99,52% Micro F1 sẽ được mô tả đúng phạm vi là text sạch từ dữ liệu tổng hợp tiếng Anh. Chất lượng trên CV thật chỉ được công bố sau lượt chạy end to end và đối chiếu với nhãn gold.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
