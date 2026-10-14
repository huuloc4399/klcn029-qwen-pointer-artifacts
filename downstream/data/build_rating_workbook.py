"""Build the Excel workbook one rater uses to score CV–JD pairs (module A).

Each rater gets an own copy with the same pairs in a rater-specific order. The workbook holds
the rubric, the JD requirements to judge, drop-down lists, live completeness checks and a
META sheet with the hashes of every input, so ratings stay traceable to the rubric and pairs
they were made against. Ratings are read back with `downstream.data.ratings_io`.

Example (run from the project root):
    python -m downstream.data.build_rating_workbook --pairs downstream/data/pairs_v1/pairs.csv \
        --jd-gold downstream/data/jd_gold_v1 --jd-text downstream/data/jd_v1 --rater A \
        --out downstream/data/labels_v1/rater_A.xlsx --lock
    python -m downstream.data.build_rating_workbook --demo   # fictional pairs, safe to share
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation

from baseline_extraction.schema import JDSchema

ROOT = Path(__file__).resolve().parents[2]
RUBRIC = ROOT / "downstream/rubric/rubric_v2.json"
DEMO_PAIRS = ROOT / "downstream/data/examples/pairs_demo.csv"
DEMO_JD_GOLD = ROOT / "downstream/data/examples/jd_gold_demo"
DEMO_OUT = ROOT / "downstream/templates/CV_JD_RATING_TEMPLATE_v1_DEMO.xlsx"
LAYOUT_VERSION = "cv_jd_rating_workbook_v1"

# (header, width, is_input). Headers are the contract with ratings_io; do not rename.
CRITERIA_COLUMNS = [
    ("pair_id", 13, False), ("cv_id", 14, False), ("jd_id", 15, False), ("criterion_id", 20, False),
    ("tieu_chi", 30, False), ("trong_so", 9, False), ("trang_thai", 16, True), ("diem", 7, True),
    ("muc_bang_chung", 16, True), ("trich_dan", 46, True), ("ly_do", 36, True), ("do_tu_tin", 12, True),
    ("ghi_chu", 26, True), ("kiem_tra", 24, False),
]
REQUIREMENT_COLUMNS = [
    ("pair_id", 13, False), ("jd_id", 13, False), ("req_id", 17, False), ("loai", 8, False),
    ("yeu_cau", 30, False), ("trang_thai", 16, True), ("muc_bang_chung", 16, True), ("trich_dan", 46, True),
    ("ghi_chu", 30, True), ("kiem_tra", 20, False),
]
OVERALL_COLUMNS = [
    ("pair_id", 13, False), ("cv_id", 14, False), ("jd_id", 15, False), ("muc_phu_hop", 12, True),
    ("nhan_xet_chung", 38, True), ("goi_y_1", 38, True), ("tieu_chi_goi_y_1", 19, True),
    ("goi_y_2", 38, True), ("tieu_chi_goi_y_2", 19, True), ("goi_y_3", 38, True),
    ("tieu_chi_goi_y_3", 19, True), ("thoi_gian_phut", 12, True), ("kiem_tra", 24, False),
]

NAVY = "17365D"
HEADER_FILL = PatternFill("solid", fgColor="D9E2EC")
LOCKED_FILL = PatternFill("solid", fgColor="F2F4F7")
BAND_FILL = PatternFill("solid", fgColor="E8EEF5")
WARN_FILL = PatternFill("solid", fgColor="FBE3B0")
OK_FONT = Font(name="Arial", size=10, color="1D7A4C", bold=True)
BASE_FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
THIN = Side(style="thin", color="C9D1DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_pairs(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = [{k.strip(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(handle)]
    missing = {"pair_id", "cv_id", "jd_id"} - set(rows[0] if rows else {})
    if missing:
        raise ValueError(f"pairs file lacks columns: {sorted(missing)}")
    seen_ids, seen_pairs = set(), set()
    for row in rows:
        if not all(row[k] for k in ("pair_id", "cv_id", "jd_id")):
            raise ValueError(f"empty id in pairs row {row}")
        if row["pair_id"] in seen_ids:
            raise ValueError(f"duplicate pair_id {row['pair_id']}")
        if (row["cv_id"], row["jd_id"]) in seen_pairs:
            raise ValueError(f"duplicate CV–JD pair {row['cv_id']} / {row['jd_id']}")
        seen_ids.add(row["pair_id"])
        seen_pairs.add((row["cv_id"], row["jd_id"]))
    return rows


def load_jd_gold(folder: Path, jd_ids: set[str]) -> tuple[dict[str, dict], str]:
    """Validated JDSchema per jd_id and one hash over the files used."""
    jds, digest = {}, hashlib.sha256()
    for jd_id in sorted(jd_ids):
        path = folder / f"{jd_id}.json"
        if not path.is_file():
            raise FileNotFoundError(f"no JD gold for {jd_id}: {path}")
        raw = path.read_bytes()
        jds[jd_id] = JDSchema.model_validate(json.loads(raw.decode("utf-8"))).model_dump()
        digest.update(jd_id.encode() + b"\0" + raw + b"\0")
        if not jds[jd_id]["requirements"]["must_have"]:
            raise ValueError(f"{jd_id} has no must_have requirement")
    return jds, digest.hexdigest()


def requirements(jd_id: str, jd: dict) -> list[tuple[str, str, str]]:
    """(req_id, type, text) in a stable order: must-have then nice-to-have."""
    out = [(f"{jd_id}-M{i:02d}", "must", text) for i, text in enumerate(jd["requirements"]["must_have"], 1)]
    out += [(f"{jd_id}-N{i:02d}", "nice", text) for i, text in enumerate(jd["requirements"]["nice_to_have"], 1)]
    return out


def ordered_pairs(pairs: list[dict], rater: str, seed: int) -> list[dict]:
    """Same pairs for every rater, shuffled per rater to spread order effects."""
    order = list(pairs)
    random.Random(f"{seed}:{rater}").shuffle(order)
    return order


# ---- sheet helpers ---------------------------------------------------------------------------

def write_table(ws, columns, rows, band_key: str | None = None) -> None:
    for c, (name, width, _) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEADER_FILL, BORDER, WRAP
        ws.column_dimensions[get_column_letter(c)].width = width
    band, previous = False, None
    for r, values in enumerate(rows, 2):
        if band_key is not None:
            key = values[0]
            if key != previous:
                band, previous = not band, key
        for c, (name, _, is_input) in enumerate(columns, 1):
            cell = ws.cell(row=r, column=c, value=values[c - 1] if c - 1 < len(values) else None)
            cell.font, cell.border, cell.alignment = BASE_FONT, BORDER, WRAP
            if is_input:
                cell.protection = Protection(locked=False)
            elif name != "kiem_tra":
                cell.fill = BAND_FILL if band else LOCKED_FILL
    ws.freeze_panes = "D2" if len(columns) > 6 else "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(2, len(rows) + 1)}"


def col(columns, name: str) -> str:
    return get_column_letter([c[0] for c in columns].index(name) + 1)


def add_list(ws, columns, name: str, source: str, rows: int, error: str) -> None:
    letter = col(columns, name)
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                        errorTitle="Giá trị không hợp lệ", error=error)
    dv.add(f"{letter}2:{letter}{rows + 1}")
    ws.add_data_validation(dv)


def add_check_format(ws, columns, rows: int) -> None:
    letter = col(columns, "kiem_tra")
    ref = f"{letter}2:{letter}{rows + 1}"
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"OK"'], font=OK_FONT))
    ws.conditional_formatting.add(ref, CellIsRule(operator="notEqual", formula=['"OK"'], fill=WARN_FILL))


def print_setup(ws, landscape: bool) -> None:
    """Print to one page width so wide sheets and long guide lines are not cut off."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1" if landscape else None


def protect(ws, allow_filter: bool = True) -> None:
    ws.protection.sheet = True
    ws.protection.autoFilter = not allow_filter  # False means "allowed" in the OOXML flag
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def text_sheet(ws, blocks: list[tuple[str, str]], width: int = 105) -> None:
    ws.column_dimensions["A"].width = width
    for r, (style, text) in enumerate(blocks, 1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            cell.font = Font(name="Arial", size=15, bold=True, color=NAVY)
        elif style == "h":
            cell.font = Font(name="Arial", size=11.5, bold=True, color=NAVY)
        elif style == "warn":
            cell.font, cell.fill = Font(name="Arial", size=10, bold=True), WARN_FILL
        else:
            cell.font = BASE_FONT


# ---- workbook --------------------------------------------------------------------------------

def guide_blocks(rubric: dict, rater: str, pair_count: int) -> list[tuple[str, str]]:
    statuses = rubric["criterion_statuses"]
    req_statuses = rubric["requirement_statuses"]
    return [
        ("title", "Hướng dẫn chấm cặp CV–JD"),
        ("p", f"Người chấm: {rater}   ·   Số cặp: {pair_count}   ·   Rubric: {rubric['rubric_version']}   ·   Mẫu: {LAYOUT_VERSION}"),
        ("warn", "Chấm độc lập. Không trao đổi điểm với người chấm khác và không xem output của mô hình trước khi nộp file."),
        ("h", "Quy trình cho mỗi cặp"),
        ("p", "1. Mở sheet CAP_CV_JD, lấy cặp theo thứ tự từ trên xuống. Mở PDF của CV và nội dung JD (sheet JD)."),
        ("p", "2. Sheet YEU_CAU_JD: với từng yêu cầu của JD, chọn trạng thái và chép nguyên văn đoạn CV làm bằng chứng."),
        ("p", "3. Sheet TIEU_CHI: chấm 8 tiêu chí theo mô tả neo ở sheet RUBRIC. Chấm C8 (phù hợp JD) sau cùng, dựa vào bước 2."),
        ("p", "4. Sheet TONG_THE: cho mức phù hợp 1–5, viết tối đa 3 gợi ý sửa CV quan trọng nhất, ghi số phút đã dùng."),
        ("p", "5. Cột kiem_tra phải hiện OK ở mọi dòng. Sheet TIEN_DO cho biết cặp nào đã xong."),
        ("h", "Quy tắc trích dẫn bằng chứng"),
        ("p", "• Chép nguyên văn từ PDF, tối đa khoảng 25 từ, không sửa chính tả. Chọn mục chứa đoạn đó ở cột muc_bang_chung."),
        ("p", "• Điểm từ 1 trở lên ở TIEU_CHI, và trạng thái met hoặc mentioned_only ở YEU_CAU_JD, bắt buộc có trích dẫn."),
        ("p", "• Với C1 và C7 (hình thức, ATS), chọn muc_bang_chung = toan_bo_cv và mô tả ngắn điều quan sát được ở cột trich_dan."),
        ("h", "Trạng thái tiêu chí (cột trang_thai ở TIEU_CHI)"),
        *[("p", f"• {key}: {text}") for key, text in statuses.items()],
        ("h", "Trạng thái yêu cầu JD (cột trang_thai ở YEU_CAU_JD)"),
        *[("p", f"• {key}: {text}") for key, text in req_statuses.items()],
        ("h", "Mức phù hợp tổng thể (cột muc_phu_hop ở TONG_THE)"),
        *[("p", f"• {text}") for text in rubric["overall_fit"]],
        ("h", "Lưu ý"),
        ("p", "• Không chấm theo tên, giới tính, ảnh, quê quán, trường học hay thông tin cá nhân. Thông tin liên hệ không dùng để cho điểm."),
        ("p", "• Không thêm, xóa hay đổi thứ tự dòng; không đổi tên sheet hoặc tiêu đề cột. Chỉ nhập ở cột trắng; cột màu xám là cố định."),
        ("p", "• File chứa trích dẫn từ CV thật: chỉ lưu trong thư mục nhóm đã thống nhất, không gửi qua kênh công khai, không đưa lên Git."),
        ("p", "• Nếu mô tả neo không đủ để quyết định, ghi lại ở cột ghi_chu để thảo luận khi phân xử."),
    ]


def build(pairs_path: Path, jd_gold_dir: Path, rater: str, out: Path, *, jd_text_dir: Path | None = None,
          rubric_path: Path = RUBRIC, seed: int = 20261013, created_at: str | None = None,
          lock: bool = False) -> dict:
    """Write one rater's workbook. With lock=True only the input columns stay editable (no password)."""
    rubric_bytes = rubric_path.read_bytes()
    rubric = json.loads(rubric_bytes.decode("utf-8"))
    criteria = rubric["criteria"]
    pairs = load_pairs(pairs_path)
    jds, jd_hash = load_jd_gold(jd_gold_dir, {p["jd_id"] for p in pairs})
    order = ordered_pairs(pairs, rater, seed)

    wb = Workbook()
    ws_guide = wb.active
    ws_guide.title = "HUONG_DAN"
    ws_rubric = wb.create_sheet("RUBRIC")
    ws_pairs = wb.create_sheet("CAP_CV_JD")
    ws_jd = wb.create_sheet("JD")
    ws_req = wb.create_sheet("YEU_CAU_JD")
    ws_crit = wb.create_sheet("TIEU_CHI")
    ws_all = wb.create_sheet("TONG_THE")
    ws_prog = wb.create_sheet("TIEN_DO")
    ws_lists = wb.create_sheet("DANH_MUC")
    ws_meta = wb.create_sheet("META")

    text_sheet(ws_guide, guide_blocks(rubric, rater, len(order)))

    # RUBRIC: one row per criterion with the five anchors.
    rubric_cols = [("criterion_id", 20, False), ("tieu_chi", 26, False), ("trong_so", 9, False),
                   ("cach_cham", 38, False)] + [(f"muc_{s}", 34, False) for s in rubric["scale"]]
    write_table(ws_rubric, rubric_cols,
                [[c["id"], c["name"], c["weight"], c["how_to_judge"], *c["anchors"]] for c in criteria])
    ws_rubric.freeze_panes = "C2"

    # CAP_CV_JD: the work list, in this rater's order. Split and intended fit are hidden on purpose.
    pair_cols = [("thu_tu", 8, False), ("pair_id", 13, False), ("cv_id", 12, False), ("jd_id", 13, False),
                 ("vi_tri", 36, False), ("cap_do", 12, False), ("so_yeu_cau_bat_buoc", 12, False),
                 ("so_yeu_cau_uu_tien", 12, False), ("tep_cv", 40, False)]
    write_table(ws_pairs, pair_cols, [
        [i, p["pair_id"], p["cv_id"], p["jd_id"], jds[p["jd_id"]]["job_title"], jds[p["jd_id"]]["seniority_level"],
         len(jds[p["jd_id"]]["requirements"]["must_have"]), len(jds[p["jd_id"]]["requirements"]["nice_to_have"]),
         p.get("cv_file", "")]
        for i, p in enumerate(order, 1)])

    # JD: requirements as the raters must judge them, plus the full text when available.
    jd_cols = [("jd_id", 13, False), ("vi_tri", 30, False), ("cap_do", 10, False), ("bat_buoc", 30, False),
               ("uu_tien", 26, False), ("trach_nhiem", 36, False), ("noi_dung_day_du", 70, False)]
    jd_rows = []
    for jd_id in sorted(jds):
        jd = jds[jd_id]
        full = ""
        if jd_text_dir and (jd_text_dir / f"{jd_id}.txt").is_file():
            full = (jd_text_dir / f"{jd_id}.txt").read_text(encoding="utf-8").strip()
        jd_rows.append([jd_id, jd["job_title"], jd["seniority_level"], "\n".join(jd["requirements"]["must_have"]),
                        "\n".join(jd["requirements"]["nice_to_have"]), "\n".join(jd["responsibilities"]),
                        full or "(chưa có tệp nội dung JD; xem các cột yêu cầu)"])
    write_table(ws_jd, jd_cols, jd_rows)

    # Lists for drop-downs.
    lists = {
        "A": ("trang_thai_tieu_chi", list(rubric["criterion_statuses"])),
        "B": ("diem", rubric["scale"]),
        "C": ("muc_bang_chung", rubric["evidence_sections"]),
        "D": ("trang_thai_yeu_cau", list(rubric["requirement_statuses"])),
        "E": ("muc_phu_hop", [1, 2, 3, 4, 5]),
        "F": ("do_tu_tin", rubric["confidence_levels"]),
        "G": ("criterion_id", [c["id"] for c in criteria]),
    }
    for letter, (title, values) in lists.items():
        ws_lists[f"{letter}1"] = title
        for i, value in enumerate(values, 2):
            ws_lists[f"{letter}{i}"] = value
    ref = {letter: f"=DANH_MUC!${letter}$2:${letter}${len(values) + 1}" for letter, (_, values) in lists.items()}

    # YEU_CAU_JD
    req_rows = [[p["pair_id"], p["jd_id"], rid, kind, text]
                for p in order for rid, kind, text in requirements(p["jd_id"], jds[p["jd_id"]])]
    write_table(ws_req, REQUIREMENT_COLUMNS, req_rows, band_key="pair_id")
    F, H, I, K = (col(REQUIREMENT_COLUMNS, n) for n in ("trang_thai", "trich_dan", "ghi_chu", "kiem_tra"))
    for r in range(2, len(req_rows) + 2):
        ws_req[f"{K}{r}"] = (f'=IF({F}{r}="","Chưa chấm",IF(OR({F}{r}="met",{F}{r}="mentioned_only"),'
                             f'IF({H}{r}="","Thiếu trích dẫn","OK"),IF(AND({F}{r}="unclear",{I}{r}=""),"Ghi lý do","OK")))')
    add_list(ws_req, REQUIREMENT_COLUMNS, "trang_thai", ref["D"], len(req_rows), "Chọn met, mentioned_only, missing hoặc unclear.")
    add_list(ws_req, REQUIREMENT_COLUMNS, "muc_bang_chung", ref["C"], len(req_rows), "Chọn một mục trong danh sách.")
    add_check_format(ws_req, REQUIREMENT_COLUMNS, len(req_rows))

    # TIEU_CHI
    crit_rows = [[p["pair_id"], p["cv_id"], p["jd_id"], c["id"], c["name"], c["weight"]] for p in order for c in criteria]
    write_table(ws_crit, CRITERIA_COLUMNS, crit_rows, band_key="pair_id")
    G, Hs, J, L, N = (col(CRITERIA_COLUMNS, n) for n in ("trang_thai", "diem", "trich_dan", "ly_do", "kiem_tra"))
    for r in range(2, len(crit_rows) + 2):
        ws_crit[f"{N}{r}"] = (
            f'=IF({G}{r}="","Chưa chấm",IF({G}{r}="scored",IF({Hs}{r}="","Thiếu điểm",'
            f'IF(AND({Hs}{r}>=1,{J}{r}=""),"Thiếu trích dẫn","OK")),'
            f'IF({G}{r}="missing_in_cv",IF(OR({Hs}{r}="",{Hs}{r}=0),"OK","missing_in_cv phải là 0"),'
            f'IF({Hs}{r}<>"","Bỏ trống điểm",IF({L}{r}="","Ghi lý do","OK")))))')
    add_list(ws_crit, CRITERIA_COLUMNS, "trang_thai", ref["A"], len(crit_rows), "Chọn scored, missing_in_cv, not_applicable hoặc unreadable.")
    add_list(ws_crit, CRITERIA_COLUMNS, "diem", ref["B"], len(crit_rows), "Điểm là số nguyên 0–4.")
    add_list(ws_crit, CRITERIA_COLUMNS, "muc_bang_chung", ref["C"], len(crit_rows), "Chọn một mục trong danh sách.")
    add_list(ws_crit, CRITERIA_COLUMNS, "do_tu_tin", ref["F"], len(crit_rows), "Chọn cao, trung_binh hoặc thap.")
    add_check_format(ws_crit, CRITERIA_COLUMNS, len(crit_rows))

    # TONG_THE
    all_rows = [[p["pair_id"], p["cv_id"], p["jd_id"]] for p in order]
    write_table(ws_all, OVERALL_COLUMNS, all_rows)
    D, Fg, Gc, Hg, Ic, Jg, Kc, Mk = (col(OVERALL_COLUMNS, n) for n in (
        "muc_phu_hop", "goi_y_1", "tieu_chi_goi_y_1", "goi_y_2", "tieu_chi_goi_y_2", "goi_y_3", "tieu_chi_goi_y_3", "kiem_tra"))
    for r in range(2, len(all_rows) + 2):
        ws_all[f"{Mk}{r}"] = (
            f'=IF({D}{r}="","Chưa chấm mức phù hợp",IF({Fg}{r}="","Cần ít nhất 1 gợi ý",'
            f'IF(OR(AND({Fg}{r}<>"",{Gc}{r}=""),AND({Hg}{r}<>"",{Ic}{r}=""),AND({Jg}{r}<>"",{Kc}{r}="")),'
            f'"Gợi ý thiếu tiêu chí","OK")))')
    add_list(ws_all, OVERALL_COLUMNS, "muc_phu_hop", ref["E"], len(all_rows), "Mức phù hợp là số nguyên 1–5.")
    for name in ("tieu_chi_goi_y_1", "tieu_chi_goi_y_2", "tieu_chi_goi_y_3"):
        add_list(ws_all, OVERALL_COLUMNS, name, ref["G"], len(all_rows), "Chọn mã tiêu chí C1–C8.")
    minutes = DataValidation(type="whole", operator="between", formula1="0", formula2="600", allow_blank=True,
                             showErrorMessage=True, error="Số phút là số nguyên 0–600.")
    letter = col(OVERALL_COLUMNS, "thoi_gian_phut")
    minutes.add(f"{letter}2:{letter}{len(all_rows) + 1}")
    ws_all.add_data_validation(minutes)
    add_check_format(ws_all, OVERALL_COLUMNS, len(all_rows))

    # TIEN_DO: live completeness per pair.
    prog_cols = [("pair_id", 13, False), ("tieu_chi_ok", 11, False), ("tong_tieu_chi", 11, False),
                 ("yeu_cau_ok", 11, False), ("tong_yeu_cau", 11, False), ("tong_the_ok", 11, False),
                 ("trang_thai", 12, False)]
    write_table(ws_prog, prog_cols, [[p["pair_id"]] for p in order])
    cid, ccheck = col(CRITERIA_COLUMNS, "pair_id"), col(CRITERIA_COLUMNS, "kiem_tra")
    rid, rcheck = col(REQUIREMENT_COLUMNS, "pair_id"), col(REQUIREMENT_COLUMNS, "kiem_tra")
    oid, ocheck = col(OVERALL_COLUMNS, "pair_id"), col(OVERALL_COLUMNS, "kiem_tra")
    for r in range(2, len(order) + 2):
        ws_prog[f"B{r}"] = f'=COUNTIFS(TIEU_CHI!${cid}:${cid},A{r},TIEU_CHI!${ccheck}:${ccheck},"OK")'
        ws_prog[f"C{r}"] = f"=COUNTIF(TIEU_CHI!${cid}:${cid},A{r})"
        ws_prog[f"D{r}"] = f'=COUNTIFS(YEU_CAU_JD!${rid}:${rid},A{r},YEU_CAU_JD!${rcheck}:${rcheck},"OK")'
        ws_prog[f"E{r}"] = f"=COUNTIF(YEU_CAU_JD!${rid}:${rid},A{r})"
        ws_prog[f"F{r}"] = f'=COUNTIFS(TONG_THE!${oid}:${oid},A{r},TONG_THE!${ocheck}:${ocheck},"OK")'
        ws_prog[f"G{r}"] = f'=IF(AND(B{r}=C{r},D{r}=E{r},F{r}=1),"Xong","Chưa")'
    ws_prog["I1"] = "Số cặp đã xong"
    ws_prog["I1"].font = BOLD
    ws_prog["I2"] = f'=COUNTIF(G2:G{len(order) + 1},"Xong")&" / {len(order)}"'
    ws_prog.column_dimensions["I"].width = 18
    ws_prog.conditional_formatting.add(f"G2:G{len(order) + 1}", CellIsRule(operator="equal", formula=['"Xong"'], font=OK_FONT))

    # META: provenance read back by ratings_io.
    meta = {
        "layout_version": LAYOUT_VERSION,
        "rater_id": rater,
        "rubric_version": rubric["rubric_version"],
        "rubric_sha256": sha256_bytes(rubric_bytes),
        "pairs_sha256": sha256_bytes(pairs_path.read_bytes()),
        "jd_gold_sha256": jd_hash,
        "order_seed": seed,
        "pair_count": len(order),
        "criterion_rows": len(crit_rows),
        "requirement_rows": len(req_rows),
        "locked": lock,
        "created_at": created_at or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }
    write_table(ws_meta, [("key", 22, False), ("value", 70, False)], [[k, v] for k, v in meta.items()])

    for ws in (ws_guide, ws_rubric, ws_pairs, ws_jd, ws_req, ws_crit, ws_all, ws_prog, ws_lists, ws_meta):
        if lock:
            protect(ws)
        print_setup(ws, landscape=ws is not ws_guide)
    ws_lists.sheet_state = "hidden"
    ws_meta.sheet_state = "hidden"
    if lock:
        wb.security = WorkbookProtection(lockStructure=True)
    wb.active = wb.sheetnames.index("HUONG_DAN")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return meta


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--pairs", type=Path)
    parser.add_argument("--jd-gold", type=Path, help="folder of <jd_id>.json files in JDSchema")
    parser.add_argument("--jd-text", type=Path, help="optional folder of <jd_id>.txt full JD texts")
    parser.add_argument("--rater", help="rater id, e.g. A or B")
    parser.add_argument("--out", type=Path)
    parser.add_argument("--seed", type=int, default=20261013)
    parser.add_argument("--lock", action="store_true",
                        help="lock fixed columns and sheets for raters (Review > Unprotect Sheet, no password)")
    parser.add_argument("--demo", action="store_true", help="build the fictional demo template")
    args = parser.parse_args()
    if args.demo:
        meta = build(DEMO_PAIRS, DEMO_JD_GOLD, "DEMO", DEMO_OUT, created_at="2026-10-13T00:00:00Z")
        print(f"{DEMO_OUT.relative_to(ROOT)}: {meta['pair_count']} pairs")
        return
    if not (args.pairs and args.jd_gold and args.rater and args.out):
        parser.error("--pairs, --jd-gold, --rater and --out are required unless --demo")
    meta = build(args.pairs, args.jd_gold, args.rater, args.out, jd_text_dir=args.jd_text, seed=args.seed,
                 lock=args.lock)
    print(json.dumps(meta, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
