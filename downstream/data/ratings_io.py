"""Read, check and export a rater's CV–JD rating workbook.

The checks repeat the workbook's kiem_tra formulas in Python and add what Excel cannot
enforce: values pasted past the drop-downs, deleted or reordered rows, and a workbook built
for another rubric. Export writes one JSON line per pair.

    python -m downstream.data.ratings_io check downstream/data/labels_v1/rater_A.xlsx
    python -m downstream.data.ratings_io export downstream/data/labels_v1/rater_A.xlsx --out rater_A.jsonl
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from downstream.data.build_rating_workbook import (
    CRITERIA_COLUMNS, LAYOUT_VERSION, OVERALL_COLUMNS, REQUIREMENT_COLUMNS, RUBRIC)


@dataclass(frozen=True)
class Issue:
    sheet: str
    row: int
    pair_id: str
    message: str

    def __str__(self) -> str:
        return f"{self.sheet}!{self.row} [{self.pair_id}] {self.message}"


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _int(value) -> int | None:
    """Whole number from a cell, or None when blank or not a whole number."""
    if value is None or _text(value) == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else None


def _rows(ws, columns) -> list[tuple[int, dict]]:
    header = [_text(c.value) for c in ws[1]][: len(columns)]
    expected = [name for name, _, _ in columns]
    if header != expected:
        raise ValueError(f"{ws.title}: header changed; expected {expected}, found {header}")
    out = []
    for index, row in enumerate(ws.iter_rows(min_row=2, max_col=len(columns), values_only=True), 2):
        if all(_text(v) == "" for v in row):
            continue
        out.append((index, dict(zip(expected, row))))
    return out


def read(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    meta = {_text(k): v for k, v in wb["META"].iter_rows(min_row=2, max_col=2, values_only=True) if k}
    return {
        "meta": meta,
        "criteria": _rows(wb["TIEU_CHI"], CRITERIA_COLUMNS),
        "requirements": _rows(wb["YEU_CAU_JD"], REQUIREMENT_COLUMNS),
        "overall": _rows(wb["TONG_THE"], OVERALL_COLUMNS),
    }


def check(book: dict, rubric_path: Path = RUBRIC) -> list[Issue]:
    rubric = json.loads(rubric_path.read_text(encoding="utf-8"))
    criterion_ids = [c["id"] for c in rubric["criteria"]]
    statuses, req_statuses = set(rubric["criterion_statuses"]), set(rubric["requirement_statuses"])
    sections, confidence = set(rubric["evidence_sections"]), set(rubric["confidence_levels"])
    meta, issues = book["meta"], []

    def add(sheet, row, pair, message):
        issues.append(Issue(sheet, row, _text(pair), message))

    if meta.get("layout_version") != LAYOUT_VERSION:
        add("META", 0, "", f"layout {meta.get('layout_version')} differs from {LAYOUT_VERSION}")
    if meta.get("rubric_version") != rubric["rubric_version"]:
        add("META", 0, "", f"built for {meta.get('rubric_version')}, current rubric is {rubric['rubric_version']}")
    for sheet, key, rows in (("TIEU_CHI", "criterion_rows", book["criteria"]),
                             ("YEU_CAU_JD", "requirement_rows", book["requirements"]),
                             ("TONG_THE", "pair_count", book["overall"])):
        if int(meta.get(key, -1)) != len(rows):
            add(sheet, 0, "", f"{len(rows)} rows, workbook was built with {meta.get(key)}; rows were added or deleted")

    per_pair: dict[str, list[str]] = {}
    for row, values in book["criteria"]:
        pair, status, score = values["pair_id"], _text(values["trang_thai"]), values["diem"]
        per_pair.setdefault(_text(pair), []).append(_text(values["criterion_id"]))
        number = _int(score)
        if status == "":
            add("TIEU_CHI", row, pair, f"{values['criterion_id']}: chưa chấm")
            continue
        if status not in statuses:
            add("TIEU_CHI", row, pair, f"trạng thái lạ '{status}'")
            continue
        if _text(score) and (number is None or not 0 <= number <= 4):
            add("TIEU_CHI", row, pair, f"điểm '{score}' không phải số nguyên 0–4")
            continue
        section = _text(values["muc_bang_chung"])
        if section and section not in sections:
            add("TIEU_CHI", row, pair, f"muc_bang_chung lạ '{section}'")
        if _text(values["do_tu_tin"]) and _text(values["do_tu_tin"]) not in confidence:
            add("TIEU_CHI", row, pair, f"do_tu_tin lạ '{values['do_tu_tin']}'")
        if status == "scored":
            if number is None:
                add("TIEU_CHI", row, pair, f"{values['criterion_id']}: thiếu điểm")
            elif number >= 1 and not _text(values["trich_dan"]):
                add("TIEU_CHI", row, pair, f"{values['criterion_id']}: điểm {number} cần trích dẫn")
        elif status == "missing_in_cv":
            if number not in (None, 0):
                add("TIEU_CHI", row, pair, f"{values['criterion_id']}: missing_in_cv phải có điểm 0")
        else:
            if number is not None:
                add("TIEU_CHI", row, pair, f"{values['criterion_id']}: {status} phải bỏ trống điểm")
            if not _text(values["ly_do"]):
                add("TIEU_CHI", row, pair, f"{values['criterion_id']}: {status} cần ghi lý do")
    for pair, ids in per_pair.items():
        if ids != criterion_ids:
            add("TIEU_CHI", 0, pair, "dòng tiêu chí bị thiếu, thừa hoặc đổi thứ tự")

    for row, values in book["requirements"]:
        pair, status = values["pair_id"], _text(values["trang_thai"])
        if status == "":
            add("YEU_CAU_JD", row, pair, f"{values['req_id']}: chưa chấm")
        elif status not in req_statuses:
            add("YEU_CAU_JD", row, pair, f"trạng thái lạ '{status}'")
        elif status in ("met", "mentioned_only") and not _text(values["trich_dan"]):
            add("YEU_CAU_JD", row, pair, f"{values['req_id']}: {status} cần trích dẫn")
        elif status == "unclear" and not _text(values["ghi_chu"]):
            add("YEU_CAU_JD", row, pair, f"{values['req_id']}: unclear cần ghi lý do")
        section = _text(values["muc_bang_chung"])
        if section and section not in sections:
            add("YEU_CAU_JD", row, pair, f"muc_bang_chung lạ '{section}'")

    for row, values in book["overall"]:
        pair, fit = values["pair_id"], _int(values["muc_phu_hop"])
        if fit is None or not 1 <= fit <= 5:
            add("TONG_THE", row, pair, "muc_phu_hop phải là số nguyên 1–5")
        if not _text(values["goi_y_1"]):
            add("TONG_THE", row, pair, "cần ít nhất 1 gợi ý")
        for k in (1, 2, 3):
            text, target = _text(values[f"goi_y_{k}"]), _text(values[f"tieu_chi_goi_y_{k}"])
            if text and target not in criterion_ids:
                add("TONG_THE", row, pair, f"goi_y_{k} cần mã tiêu chí hợp lệ")
        if _text(values["thoi_gian_phut"]) and _int(values["thoi_gian_phut"]) is None:
            add("TONG_THE", row, pair, "thoi_gian_phut phải là số nguyên")
    return issues


def to_records(book: dict) -> list[dict]:
    meta = book["meta"]
    records: dict[str, dict] = {}
    for _, v in book["overall"]:
        pair = _text(v["pair_id"])
        records[pair] = {
            "pair_id": pair, "cv_id": _text(v["cv_id"]), "jd_id": _text(v["jd_id"]),
            "rater_id": _text(meta.get("rater_id")), "rubric_version": _text(meta.get("rubric_version")),
            "rubric_sha256": _text(meta.get("rubric_sha256")), "pairs_sha256": _text(meta.get("pairs_sha256")),
            "overall_fit": _int(v["muc_phu_hop"]), "comment": _text(v["nhan_xet_chung"]),
            "suggestions": [{"text": _text(v[f"goi_y_{k}"]), "criterion_id": _text(v[f"tieu_chi_goi_y_{k}"])}
                            for k in (1, 2, 3) if _text(v[f"goi_y_{k}"])],
            "minutes": _int(v["thoi_gian_phut"]),
            "criteria": {}, "requirements": {},
        }
    for _, v in book["criteria"]:
        status = _text(v["trang_thai"])
        score = _int(v["diem"])
        records[_text(v["pair_id"])]["criteria"][_text(v["criterion_id"])] = {
            "status": status, "score": 0 if status == "missing_in_cv" else score,
            "evidence_section": _text(v["muc_bang_chung"]), "quote": _text(v["trich_dan"]),
            "reason": _text(v["ly_do"]), "confidence": _text(v["do_tu_tin"]), "note": _text(v["ghi_chu"]),
        }
    for _, v in book["requirements"]:
        records[_text(v["pair_id"])]["requirements"][_text(v["req_id"])] = {
            "type": _text(v["loai"]), "text": _text(v["yeu_cau"]), "status": _text(v["trang_thai"]),
            "evidence_section": _text(v["muc_bang_chung"]), "quote": _text(v["trich_dan"]), "note": _text(v["ghi_chu"]),
        }
    return [records[k] for k in sorted(records)]


def main() -> int:
    parser = argparse.ArgumentParser(description="Check or export a CV–JD rating workbook.")
    parser.add_argument("command", choices=("check", "export"))
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out", type=Path, help="JSONL output for export")
    parser.add_argument("--allow-incomplete", action="store_true", help="export only pairs without issues")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    book = read(args.workbook)
    issues = check(book)
    pairs = [_text(v["pair_id"]) for _, v in book["overall"]]
    broken = {i.pair_id for i in issues}
    workbook_level = any(i.pair_id == "" for i in issues)
    done = [p for p in pairs if p not in broken]
    print(f"{args.workbook.name}: rater {book['meta'].get('rater_id')}, {len(done)}/{len(pairs)} pairs complete, {len(issues)} issues")
    for issue in issues[:60]:
        print(f"  {issue}")
    if len(issues) > 60:
        print(f"  … {len(issues) - 60} more")
    if args.command == "check":
        return 0 if not issues else 1
    if not args.out:
        parser.error("export needs --out")
    if workbook_level or (issues and not args.allow_incomplete):
        print("Not exported: fix the issues above, or pass --allow-incomplete to export complete pairs only.")
        return 1
    records = [r for r in to_records(book) if r["pair_id"] in done]
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in records), encoding="utf-8")
    print(f"Exported {len(records)} pairs to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
