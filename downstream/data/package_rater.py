"""Zip everything one rater needs: the locked workbook, the CV PDFs of their pairs and a guide.

The package holds real CVs, so it is written under downstream/data/labels_v1/packages/
(gitignored) and must only travel through the team's private channel. It deliberately leaves
out pairs.csv (split, intended fit, CV track), JD gold files, CVSchema gold and model outputs.

    python -m downstream.data.package_rater --rater A --split dev     # pilot round
    python -m downstream.data.package_rater --rater B --split dev
    python -m downstream.data.package_rater --rater A                 # full round, all 48 pairs
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from downstream.data import build_rating_workbook as builder

ROOT = Path(__file__).resolve().parents[2]
PAIRS = ROOT / "downstream/data/pairs_v1/pairs.csv"
JD_GOLD = ROOT / "downstream/data/jd_gold_v1"
JD_TEXT = ROOT / "downstream/data/jd_v1"
CV_MANIFEST = ROOT / "baseline_extraction/data/manifest.csv"
OUT_DIR = ROOT / "downstream/data/labels_v1/packages"
NAVY = RGBColor(0x17, 0x36, 0x5D)
# Bump when the guide's rules change; raters quote it so the coordinator knows which rules they used.
GUIDE_VERSION = "2026-10-17, bản 3 (sau phân xử vòng chấm thử)"


def cv_pdf_paths(cv_ids: set[str], manifest: Path = CV_MANIFEST) -> dict[str, Path]:
    with manifest.open(encoding="utf-8-sig", newline="") as handle:
        rows = {r["id"]: ROOT / r["path"] for r in csv.DictReader(handle)}
    missing = sorted(c for c in cv_ids if c not in rows or not rows[c].is_file())
    if missing:
        raise FileNotFoundError(f"no PDF for {missing}")
    return {c: rows[c] for c in cv_ids}


def guide(path: Path, *, rater: str, round_name: str, workbook_name: str, meta: dict, cv_ids: list[str],
          deadline: str, coordinator: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10.5)

    def heading(text: str, size: float = 13) -> None:
        run = doc.add_paragraph().add_run(text)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(size), NAVY

    def bullets(items: list[str]) -> None:
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    def steps(items: list[str]) -> None:
        for item in items:
            doc.add_paragraph(item, style="List Number")

    def checklist(items: list[str]) -> None:
        """Numbered (1), (2)… restarting for each list, so the level rules can cite them."""
        for number, item in enumerate(items, 1):
            doc.add_paragraph(f"({number}) {item}", style="List Bullet")

    heading(f"Hướng dẫn chấm cặp CV–JD · {round_name} · Người chấm {rater}", 15)
    doc.add_paragraph(f"Số cặp: {meta['pair_count']} · Số CV: {len(cv_ids)} · Rubric: {meta['rubric_version']} · "
                      f"Hướng dẫn bản {GUIDE_VERSION} · Hạn nộp: {deadline} · Liên hệ khi file lỗi: {coordinator}")

    heading("1. Mục đích")
    doc.add_paragraph(
        "Bạn đánh giá mức phù hợp giữa từng CV và một mô tả công việc (JD) theo rubric 8 tiêu chí. Kết quả là "
        "nhãn tham chiếu của con người để so sánh với điểm do hệ thống chấm. Hai người chấm độc lập cùng các cặp; "
        "chỗ khác nhau sẽ được phân xử sau.")

    heading("2. Trong gói có gì")
    bullets([
        f"{workbook_name}: file Excel để chấm. Chỉ sửa ở các cột trắng; cột xám đã khóa.",
        f"CV_PDF/: {len(cv_ids)} CV đặt tên theo mã (ví dụ EN002.pdf). Mã trong file Excel trùng với tên tệp.",
        "Nội dung JD nằm ngay trong sheet JD của file Excel.",
        "Tệp này (00_DOC_TRUOC.docx).",
    ])

    heading("3. Quy tắc bắt buộc")
    bullets([
        "Chấm độc lập: không trao đổi điểm hay cách chấm với người chấm khác cho tới khi cả hai đã nộp.",
        "Không dùng ChatGPT, Claude hay công cụ AI nào để đọc CV hoặc gợi ý điểm, và không xem kết quả của mô hình.",
        "CV là dữ liệu thật của sinh viên: chỉ mở trên máy của bạn, không gửi lên nhóm chat chung, mạng xã hội "
        "hay công cụ online. Xóa gói khỏi máy sau khi nhóm xác nhận đã nhận file chấm.",
        "Không chấm theo tên, giới tính, ảnh, quê quán, trường học hay thông tin cá nhân.",
        "Không thêm, xóa, đổi thứ tự dòng; không đổi tên sheet, tên cột hay tên file.",
        "Mở bằng Microsoft Excel (hoặc LibreOffice). Không mở bằng Google Sheets vì có thể mất danh sách thả xuống.",
    ])

    heading("4. Cách chấm một cặp")
    steps([
        "Sheet CAP_CV_JD: lấy cặp tiếp theo theo cột thu_tu. Mở CV_PDF/<cv_id>.pdf và đọc JD ở sheet JD.",
        "Sheet YEU_CAU_JD: với mỗi yêu cầu của JD, chọn trang_thai: met (có dùng trong dự án, kinh nghiệm, chứng chỉ), "
        "mentioned_only (chỉ liệt kê ở mục kỹ năng/mục tiêu), missing (không thấy), unclear (không đủ để quyết định, "
        "ghi lý do ở ghi_chu). Với met và mentioned_only, chép nguyên văn đoạn CV vào trich_dan.",
        "Sheet TIEU_CHI: chấm 0–4 cho 8 tiêu chí theo mô tả ở sheet RUBRIC. Chấm C8 (phù hợp JD) sau cùng, dựa vào bước 2.",
        "Sheet TONG_THE: chọn muc_phu_hop 1–5, viết tối đa 3 gợi ý sửa CV quan trọng nhất (mỗi gợi ý chọn một tiêu chí), "
        "ghi số phút đã dùng cho cặp.",
        "Kiểm tra cột kiem_tra hiện OK ở mọi dòng của cặp. Sheet TIEN_DO cho biết cặp đã xong chưa.",
    ])

    heading("Một số trường hợp khó ở YEU_CAU_JD", 11.5)
    bullets([
        "Bỏ qua mức độ: JD ghi 'thành thạo Java', CV ghi 'Java' trong mục kỹ năng → mentioned_only như bình thường. "
        "Mức độ thành thạo được phản ánh ở tiêu chí C3.",
        "CV nêu một khả năng khác cùng lĩnh vực → missing và ghi lý do ở ghi_chu. Ví dụ JD cần giao tiếp tiếng Anh "
        "(viết và nói), CV chỉ ghi đọc hiểu tài liệu kỹ thuật.",
        "Yêu cầu có lựa chọn thay thế ('Jest hoặc Jasmine', 'React hoặc Vue') → CV có một trong các lựa chọn là đủ.",
        "Bao hàm trực tiếp (dùng công nghệ trong CV thì bắt buộc dùng kỹ năng được yêu cầu) → xem như CV có kỹ năng đó: "
        "trong dự án/kinh nghiệm → met, chỉ ở mục kỹ năng → mentioned_only; ghi ở ghi_chu đã suy ra từ đâu. Gồm: "
        "Spring Boot → Java; Laravel → PHP; Django, FastAPI, Flask → Python; Express, NestJS → Node.js; React, Vue, "
        "Angular → JavaScript; ASP.NET → C#; dựng backend/API bằng FastAPI, Express, Spring Boot, Django → REST API. "
        "Ví dụ: 'Integrated a PhoBERT-based chatbot through a FastAPI backend' → REST API: met, ghi chú 'suy ra từ FastAPI backend'.",
        "Liên quan lỏng (thường đi cùng nhưng không bắt buộc) → unclear và ghi lý do. Ví dụ: Docker → Linux; MySQL → thiết kế "
        "hay tối ưu CSDL; Git → CI/CD; React → TypeScript; Java, C# → OOP (trừ khi CV nêu OOP, class, kế thừa, design pattern).",
        "Câu tự nhận kinh nghiệm ('Experienced in X', 'Có kinh nghiệm với X') ở mục kỹ năng hoặc giới thiệu mà không gắn với "
        "dự án cụ thể → mentioned_only. Nếu X xuất hiện trong một dự án, trích đoạn dự án đó → met.",
        "Ngoại ngữ và kỹ năng mềm: có chứng chỉ (TOEIC, IELTS) hoặc việc làm cụ thể (thuyết trình bằng tiếng Anh, làm trưởng "
        "nhóm) → met; chỉ tự ghi ở mục kỹ năng hoặc giới thiệu → mentioned_only. CV viết bằng tiếng Anh không phải bằng chứng "
        "về tiếng Anh → unclear, ghi chú 'chỉ suy ra từ ngôn ngữ của CV'.",
        "Kỹ năng mềm chỉ tính khi đúng tên hoặc đồng nghĩa: tư duy logic = logical thinking; giải quyết vấn đề = problem-solving; "
        "tư duy phân tích = analytical thinking; làm việc nhóm = teamwork; giao tiếp = communication. Kỹ năng khác không thay "
        "thế nhau (problem-solving không tính cho tư duy logic) → missing.",
        "Công cụ cụ thể nhưng CV dùng công cụ khác cùng loại (GitHub thay GitLab, MySQL thay PostgreSQL, Trello thay Jira) → "
        "missing, ghi chú 'dùng X thay vì Y'. Chỉ tính thay thế khi JD đã ghi 'hoặc', 'or similar', 'tương đương'.",
        "Yêu cầu là một loại công cụ kèm ví dụ ('E-commerce platforms: Shopify, WordPress or other') → chỉ tính công cụ cùng loại "
        "(WooCommerce, Magento…); tự xây một website bán hàng không tính → missing, ghi chú.",
        "Yêu cầu về lĩnh vực (tài chính – ngân hàng, fintech) → dự án phải thuộc lĩnh vực đó. Chỉ có một tính năng liên quan "
        "trong dự án khác lĩnh vực (tích hợp cổng thanh toán trong web bán hàng) → unclear, ghi chú.",
        "Yêu cầu là một nhóm kỹ thuật kèm ví dụ trong ngoặc ('Hiệu năng web (critical rendering path, code splitting, "
        "lazy loading)') → met khi dự án hoặc kinh nghiệm nêu ít nhất một kỹ thuật cụ thể trong ngoặc hoặc cùng loại "
        "(lazy loading, memoization, nén ảnh); chỉ ghi chung chung 'tối ưu hiệu năng' → unclear, ghi chú; không có → missing.",
        "Trước khi chọn missing: dùng Ctrl+F trong PDF tìm tên công nghệ, cả tên viết tắt và cách viết khác (TS, "
        "TypeScript; Postgres, PostgreSQL; JS, JavaScript). Ở vòng chấm thử, phần lớn các ca một người chọn met còn "
        "người kia chọn missing là do bỏ sót khi đọc.",
        "Với missing: để trống muc_bang_chung và trich_dan; lý do (nếu cần) ghi ở ghi_chu.",
    ])

    heading("Yêu cầu khái niệm và khi nào chọn unclear (bổ sung bản 3)", 11.5)
    doc.add_paragraph(
        "Ở vòng chấm thử, không ai dùng unclear, và phần lớn bất đồng nằm ở các yêu cầu là khái niệm hoặc kỹ thuật "
        "(Responsive, OOP, thiết kế giao diện) chứ không phải tên công nghệ. Các quy tắc dưới đây được chốt khi phân xử "
        "vòng thử; chúng cụ thể hóa quy tắc 'liên quan lỏng → unclear' ở trên.")
    bullets([
        "Yêu cầu là khái niệm hoặc kỹ thuật cần CV nêu chính khái niệm đó, hoặc một kỹ thuật cụ thể thuộc nó. Ví dụ "
        "Responsive: responsive, mobile-first, media query, giao diện tương thích di động. Có trong dự án → met; chỉ ở mục "
        "kỹ năng → mentioned_only.",
        "Chỉ có công nghệ nền mà không nêu khái niệm → unclear, ghi chú 'chỉ có X'. Ví dụ: CSS → Responsive; HTML → "
        "Semantic/SEO-friendly HTML; CSS → Modern CSS (Tailwind, Flexbox, Grid); HTML, CSS → thiết kế giao diện web; "
        "thư viện biểu đồ (Chart.js) → Responsive.",
        "Template hoặc framework giao diện (Thymeleaf, Blade, Bootstrap) → HTML/CSS: liên quan lỏng → unclear. HTML/CSS chỉ "
        "met hoặc mentioned_only khi CV ghi HTML, CSS (hoặc HTML5, CSS3).",
        "Tích hợp API của mô hình ngôn ngữ (Gemini, OpenAI) → LLM prompting: unclear. Met khi CV nêu thiết kế prompt, "
        "prompt engineering hoặc kỹ thuật tương đương.",
        "OOP: CV nêu OOP, class, kế thừa, design pattern trong dự án → met. Chỉ có trong danh sách môn học hoặc mục kỹ năng → "
        "mentioned_only. Chỉ dùng Java hoặc C# → unclear (đã có ở trên).",
        "Kiểm thử: 'test' chung chung, stress test, testing mindset không phải Unit testing hay Integration testing → missing "
        "(khả năng khác cùng lĩnh vực).",
        "Yêu cầu lĩnh vực hoặc loại dự án (thương mại điện tử, trang quản trị, ERP, tài chính): dự án thuộc đúng loại đó → met; "
        "dự án chỉ gần loại đó (ứng dụng quản lý nhỏ khi JD cần trang quản trị/ERP) → unclear.",
        "Kỹ năng mềm: chỉ tính khi đúng tên hoặc đồng nghĩa (đã có ở trên). CV ghi kỹ năng đó và có việc làm cụ thể tương ứng "
        "trong dự án hoặc kinh nghiệm → met. Chỉ có việc làm mà CV không nêu tên kỹ năng ('xử lý logic game' với yêu cầu "
        "tư duy logic) → missing.",
        "Tóm lại: met = có bằng chứng sử dụng; mentioned_only = CV nêu đúng yêu cầu nhưng chỉ liệt kê; unclear = CV có thứ "
        "liên quan nhưng không đủ để khẳng định; missing = không có gì liên quan. Phân vân giữa unclear và missing: CV có "
        "công nghệ hay việc làm gần với yêu cầu thì chọn unclear và ghi rõ là gì.",
    ])

    heading("5. Chấm tiêu chí")
    bullets([
        "scored: chấm 0–4 theo mô tả neo. Điểm từ 1 trở lên phải có trích dẫn.",
        "missing_in_cv: PDF đọc được nhưng CV không có nội dung cho tiêu chí đó. Điểm bắt buộc là 0.",
        "not_applicable: tiêu chí không áp dụng cho vị trí. Bỏ trống điểm, ghi lý do.",
        "unreadable: không đọc chắc được phần liên quan (mờ, vỡ font, mất trang). Bỏ trống điểm, ghi lý do.",
        "C1 (trình bày) và C7 (ATS): chọn muc_bang_chung = toan_bo_cv và mô tả ngắn điều quan sát được. Với C7, thử bôi đen "
        "và sao chép chữ trong PDF.",
        "C5 (kinh nghiệm) với JD intern/fresher: nếu không có kinh nghiệm nhưng có dự án đạt mức 3 ở C4, cho tối đa 2 điểm "
        "và ghi 'dự án thay thế' ở ly_do.",
        "Mức 0 vì CV không có mục hoặc nội dung cho tiêu chí (không có kinh nghiệm, không có chứng chỉ…) → luôn dùng "
        "missing_in_cv, kể cả C5 với JD junior. Chỉ dùng scored với điểm 0 khi CV có nội dung nhưng yếu đến mức 0 "
        "(ví dụ C8 dưới 20%).",
        "C8: tính tỉ lệ = (số met + 0,5 × số mentioned_only) / số yêu cầu bắt buộc, không tính các dòng unclear. Làm đúng "
        "theo khoảng: dưới 0,2 → 0; 0,2 đến dưới 0,4 → 1; 0,4 đến dưới 0,6 → 2; 0,6 đến dưới 0,8 → 3; từ 0,8 → 4 "
        "(cần thêm vài yêu cầu ưu tiên). Ví dụ 0,59 → 2.",
        "C8 phải đúng bằng mức tính từ chính các trạng thái bạn chọn ở YEU_CAU_JD. Ở vòng thử, C8 nhiều lần không khớp với "
        "các trạng thái đó. Khi lập nhãn cuối cùng, C8 được tính lại từ trạng thái yêu cầu sau phân xử.",
    ])

    heading("Cách chọn mức điểm (bổ sung sau vòng chấm thử)", 11.5)
    doc.add_paragraph(
        "Ở vòng chấm thử, mức 4 được dùng cho phần lớn CV ở C1, C2, C3 và C7. Khi gần như CV nào cũng được 4, điểm "
        "không còn phân biệt được CV tốt với CV trung bình. Các quy tắc dưới đây không đổi mô tả mức trong sheet "
        "RUBRIC; chúng chỉ nói cách kiểm tra từng điều kiện.")
    bullets([
        "Chấm từ dưới lên: đọc mô tả mức 2; CV thỏa hết điều kiện của mức 2 thì xét mức 3, rồi mức 4. Dừng ở mức cao "
        "nhất mà CV thỏa mọi điều kiện. Thiếu một điều kiện là chưa đạt mức đó.",
        "Mức 4 nghĩa là theo tiêu chí này CV gần như không còn gì để sửa. Khi cho 4, ghi ở ly_do vì sao CV đạt từng "
        "điều kiện của mức 4.",
        "Thấy mình sắp cho 4 ở một tiêu chí cho hầu hết CV: dừng lại, đọc lại mô tả mức 4 và các danh sách kiểm tra dưới đây.",
    ])

    heading("C1 · Cấu trúc và hình thức (chỉ xem cách trình bày, không xem nội dung chuyên môn hay máy đọc)", 11)
    checklist([
        "Có đủ các mục chính không: mục tiêu hoặc giới thiệu, kỹ năng, dự án hoặc kinh nghiệm, học vấn?",
        "Định dạng có nhất quán không: cách ghi ngày tháng, gạch đầu dòng, cỡ và kiểu chữ giữa các mục?",
        "Đếm lỗi chính tả và lỗi khoảng trắng hoặc dòng tràn sang trang mới.",
        "Độ dài 1–2 trang, không dày đặc chữ?",
        "Thông tin quan trọng với vị trí (kỹ năng, dự án hoặc kinh nghiệm CNTT) có đặt trước thông tin ít liên quan "
        "(sở thích, thông tin phụ) không?",
    ])
    bullets([
        "Mức 4: đạt cả (1)–(5) và không có lỗi chính tả nào. Mức 3: đạt (1), (2), (4) và chỉ có lỗi nhỏ (vài lỗi chính tả, "
        "khoảng trắng, một dòng tràn trang). Mức 2: đủ mục nhưng (2) hoặc (4) không đạt. Mức 1 và 0 theo sheet RUBRIC.",
        "Ghi vào ly_do kết quả từng câu, ví dụ: 'đủ mục; ngày tháng lúc ghi 03/2025 lúc ghi Mar 2025; 2 lỗi chính tả → 2'.",
    ])

    heading("C2 · Nội dung chuyên môn CNTT (chấm độc lập với JD)", 11)
    checklist([
        "Mục tiêu hoặc giới thiệu có nêu một hướng cụ thể (Backend, Frontend, Mobile…) không?",
        "Dự án và kinh nghiệm có cùng hướng đó không?",
        "Dự án hoặc kinh nghiệm có ít nhất một dấu hiệu chiều sâu kỹ thuật: thiết kế (CSDL, API, kiến trúc), kiểm thử "
        "(unit test, integration test) hoặc triển khai (deploy, Docker, CI/CD)? Chỉ ghi ở mục kỹ năng thì không tính.",
    ])
    bullets([
        "Mức 4: đạt cả (1)–(3). Mức 3: đạt (1) và (2). Mức 2: có nội dung CNTT nhưng không đạt (1), hoặc dự án lệch hướng đã nêu.",
    ])

    heading("C3 · Kỹ năng chuyên môn và bổ trợ", 11)
    checklist([
        "Lấy tối đa 5 kỹ năng chuyên môn chính (đầu danh sách, hoặc những kỹ năng CV nhấn mạnh). Đếm bao nhiêu kỹ năng "
        "xuất hiện trong dự án, kinh nghiệm hoặc chứng chỉ.",
        "Danh sách kỹ năng có được nhóm rõ không (ngôn ngữ, framework, CSDL, công cụ…)?",
        "Có kỹ năng bổ trợ không? Có ngữ cảnh cụ thể không: vai trò trong nhóm ('trưởng nhóm 4 người'), thuyết trình đồ "
        "án, chứng chỉ ngoại ngữ có điểm?",
    ])
    bullets([
        "Mức 4: ít nhất 4/5 kỹ năng có minh chứng, danh sách được nhóm, và kỹ năng bổ trợ có ngữ cảnh. Mức 3: ít nhất 3/5 "
        "có minh chứng và có kỹ năng bổ trợ (chưa cần ngữ cảnh). Mức 2: dưới 3/5 có minh chứng, hoặc không có kỹ năng bổ trợ.",
        "Ghi số đếm vào ly_do, ví dụ: '3/5 có minh chứng (Java, Spring Boot, MySQL); soft skill chỉ liệt kê → 3'.",
    ])

    heading("C7 · Khả năng đáp ứng ATS (máy đọc được): kiểm tra bắt buộc", 11)
    checklist([
        "Mở PDF, nhấn Ctrl+A rồi Ctrl+C, dán vào Notepad (không dán vào công cụ online). Có chép ra được chữ không? "
        "Không, hoặc gần như không → 0.",
        "Chữ dán ra có đúng thứ tự đọc không, hay các mục bị trộn, cột trái và cột phải xen nhau từng dòng? Bị trộn ở "
        "nhiều chỗ, hoặc thông tin quan trọng (kỹ năng, liên hệ, ngày tháng) chỉ nằm trong hình hay biểu tượng → 1.",
        "Đếm các vấn đề còn lại: (a) bố cục hai cột; (b) thanh, sao hay vòng tròn đo kỹ năng bằng đồ họa; (c) tiêu đề "
        "mục lạ, không phải Mục tiêu, Kỹ năng, Kinh nghiệm, Dự án, Học vấn hay tương đương tiếng Anh; (d) ngày tháng "
        "không nhất quán; (e) biểu tượng thay chữ (icon điện thoại không kèm số).",
    ])
    bullets([
        "Mức 4: chép đúng thứ tự và không có vấn đề nào ở (3). Mức 3: chỉ có hai cột đơn giản mà chữ dán ra vẫn đúng "
        "thứ tự, hoặc chỉ một lỗi nhỏ. Mức 2: còn một vấn đề rõ (thanh đồ họa kỹ năng, tiêu đề lạ, hai cột làm thứ tự "
        "dán ra bị xáo một phần). Mức 1: nhiều vấn đề, hoặc theo (2).",
        "Ghi vào ly_do kết quả (1)–(3), ví dụ: 'chép được; thứ tự đúng; có thanh đồ họa kỹ năng → 2'.",
    ])

    heading("6. Trích dẫn")
    bullets([
        "Chép nguyên văn từ PDF, tối đa khoảng 25 từ, không sửa chính tả; chọn mục chứa đoạn đó ở muc_bang_chung.",
        "Trích một đoạn liền mạch; không nối nhiều chỗ bằng '...'. Kỹ năng xuất hiện ở nhiều nơi thì trích chỗ mạnh nhất "
        "(dự án hoặc kinh nghiệm trước, mục kỹ năng sau).",
        "Mục trong CV có tên khác danh sách thì chọn mục tương ứng: Ngoại ngữ → skills (chứng chỉ có điểm như TOEIC, "
        "IELTS → certifications); Mục tiêu, Giới thiệu → summary; Thực tập, Công việc → experience; Đồ án → projects; "
        "Giải thưởng, Thành tích → awards; Hoạt động, CLB, Tình nguyện → activities. Chỉ dùng toan_bo_cv cho C1 và C7.",
        "Chỉ chấm theo nội dung CV; không mở link GitHub, portfolio hay demo. Link chỉ tính như một dòng chữ trong CV "
        "(với C4, link repo hoặc demo của dự án vẫn là 'sản phẩm kiểm chứng được').",
        "Link GitHub với yêu cầu Git: link gắn với một dự án cụ thể → met (projects); chỉ có link hồ sơ ở phần liên hệ → "
        "mentioned_only (personal_info); Git/GitHub chỉ ghi trong mục kỹ năng → mentioned_only (skills).",
        "Ví dụ: tiêu chí C4, điểm 3, muc_bang_chung = projects, trich_dan = 'Built REST API for booking with Spring Boot "
        "and MySQL, role: backend developer', ly_do = 'Có vai trò, công nghệ, đầu việc; chưa có kết quả đo được'.",
        "Không chắc giữa hai mức điểm: chọn mức thấp hơn khi thiếu bằng chứng, và ghi điều phân vân vào ghi_chu. "
        "Ghi chú này rất có ích cho buổi chỉnh rubric.",
    ])

    heading("7. Thời gian và nộp bài")
    bullets([
        "Khoảng 10–12 phút mỗi cặp; cặp đầu tiên lâu hơn. Có thể chia nhiều buổi, nhớ lưu file thường xuyên (Ctrl+S).",
        f"Hạn nộp: {deadline}. Gửi lại đúng file {workbook_name} qua kênh riêng của nhóm cho {coordinator}.",
        "Không cần làm hết trong một lần; file chấm dở vẫn kiểm tra được.",
    ])

    heading("Thông tin truy vết", 10)
    doc.add_paragraph(f"layout {meta['layout_version']} · rubric_sha256 {meta['rubric_sha256'][:16]} · "
                      f"pairs_sha256 {meta['pairs_sha256'][:16]} · jd_gold_sha256 {meta['jd_gold_sha256'][:16]} · "
                      f"split {meta['split']} · tạo lúc {meta['created_at']}")
    doc.save(path)


def package(rater: str, *, split: str | None = None, pairs: Path = PAIRS, jd_gold: Path = JD_GOLD,
            jd_text: Path = JD_TEXT, out_dir: Path = OUT_DIR, deadline: str = "do điều phối viên thông báo",
            coordinator: str = "điều phối viên của nhóm", pdf_paths=cv_pdf_paths, guide_only: bool = False) -> Path:
    """Write the rater's zip; with guide_only, write only <stem>_00_DOC_TRUOC.docx (rule changes in the guide)."""
    round_name = "Vòng chấm thử" if split == "dev" else "Vòng chính" if split in (None, "test") else f"Tập {split}"
    stem = f"{ {'dev': 'pilot', None: 'full', 'test': 'main'}.get(split, split)}_rater_{rater}"
    workbook_name = f"{stem}.xlsx"
    with tempfile.TemporaryDirectory() as folder:
        work = Path(folder)
        meta = builder.build(pairs, jd_gold, rater, work / workbook_name, jd_text_dir=jd_text, lock=True, split=split)
        with pairs.open(encoding="utf-8-sig", newline="") as handle:
            cv_ids = sorted({r["cv_id"] for r in csv.DictReader(handle) if split is None or r.get("split") == split})
        guide(work / "00_DOC_TRUOC.docx", rater=rater, round_name=round_name, workbook_name=workbook_name, meta=meta,
              cv_ids=cv_ids, deadline=deadline, coordinator=coordinator)
        out_dir.mkdir(parents=True, exist_ok=True)
        if guide_only:
            target = out_dir / f"{stem}_00_DOC_TRUOC.docx"
            target.write_bytes((work / "00_DOC_TRUOC.docx").read_bytes())
            return target
        pdfs = pdf_paths(set(cv_ids))
        target = out_dir / f"{stem}.zip"
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.write(work / "00_DOC_TRUOC.docx", f"{stem}/00_DOC_TRUOC.docx")
            archive.write(work / workbook_name, f"{stem}/{workbook_name}")
            for cv in cv_ids:
                archive.write(pdfs[cv], f"{stem}/CV_PDF/{cv}.pdf")
    record = {"package": target.name, "rater": rater, "split": split or "all", "pairs": meta["pair_count"], "cvs": len(cv_ids),
              "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
              "created_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}
    with (out_dir / "PACKAGES_LOG.jsonl").open("a", encoding="utf-8") as log:
        log.write(json.dumps(record, ensure_ascii=False) + "\n")
    return target


def main() -> None:
    parser = argparse.ArgumentParser(description="Build one rater's package.")
    parser.add_argument("--rater", required=True)
    parser.add_argument("--split", help="dev for the pilot round; omit for all pairs")
    parser.add_argument("--deadline", default="do điều phối viên thông báo")
    parser.add_argument("--coordinator", default="điều phối viên của nhóm")
    parser.add_argument("--guide-only", action="store_true",
                        help="only write the guide .docx; use when just the rules in the guide changed")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    target = package(args.rater, split=args.split, deadline=args.deadline, coordinator=args.coordinator,
                     guide_only=args.guide_only)
    print(f"{target.relative_to(ROOT)} ({target.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
