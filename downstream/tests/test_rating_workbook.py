"""Module A rating workbook: build, fill like a rater would, check and export.

Run from the project root: python -m unittest downstream/tests/test_rating_workbook.py -v
"""

from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from downstream.data import build_rating_workbook as builder
from downstream.data import ratings_io

RUBRIC = json.loads(builder.RUBRIC.read_text(encoding="utf-8"))
CRITERIA = [c["id"] for c in RUBRIC["criteria"]]


def column(ws, name: str) -> int:
    return [c.value for c in ws[1]].index(name) + 1


def fill_complete(path: Path) -> None:
    """Fill every input the way a careful rater would."""
    wb = load_workbook(path)
    ws = wb["TIEU_CHI"]
    status, score, quote, section = (column(ws, n) for n in ("trang_thai", "diem", "trich_dan", "muc_bang_chung"))
    for r in range(2, ws.max_row + 1):
        crit = ws.cell(r, column(ws, "criterion_id")).value
        if crit == "C6_certs_activities":
            ws.cell(r, status, "missing_in_cv")
            ws.cell(r, score, 0)
        else:
            ws.cell(r, status, "scored")
            ws.cell(r, score, 3)
            ws.cell(r, section, "projects")
            ws.cell(r, quote, "Built REST API with Spring Boot")
    ws = wb["YEU_CAU_JD"]
    for r in range(2, ws.max_row + 1):
        met = ws.cell(r, column(ws, "loai")).value == "must"
        ws.cell(r, column(ws, "trang_thai"), "met" if met else "missing")
        if met:
            ws.cell(r, column(ws, "trich_dan"), "Spring Boot, MySQL")
    ws = wb["TONG_THE"]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, column(ws, "muc_phu_hop"), 4)
        ws.cell(r, column(ws, "goi_y_1"), "Thêm kết quả đo được cho dự án chính.")
        ws.cell(r, column(ws, "tieu_chi_goi_y_1"), "C4_projects")
        ws.cell(r, column(ws, "thoi_gian_phut"), 9)
    wb.save(path)


class RatingWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)

    def build(self, rater="A", pairs=builder.DEMO_PAIRS, jd=builder.DEMO_JD_GOLD, lock=False):
        out = self.tmp / f"rater_{rater}.xlsx"
        meta = builder.build(pairs, jd, rater, out, created_at="2026-10-13T00:00:00Z", lock=lock)
        return out, meta

    def test_layout_and_traceability(self):
        path, meta = self.build()
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["HUONG_DAN", "RUBRIC", "CAP_CV_JD", "JD", "YEU_CAU_JD", "TIEU_CHI",
                                         "TONG_THE", "TIEN_DO", "DANH_MUC", "META"])
        self.assertEqual((meta["pair_count"], meta["criterion_rows"]), (3, 3 * len(CRITERIA)))
        # BE has 5 must + 3 nice; FE has 4 + 2 and is used twice.
        self.assertEqual(meta["requirement_rows"], 8 + 6 + 6)
        self.assertFalse(wb.security and wb.security.lockStructure)
        self.assertFalse(wb["RUBRIC"].protection.sheet)
        headers = [c.value for c in wb["CAP_CV_JD"][1]]
        self.assertNotIn("intended_fit", headers)
        self.assertNotIn("split", headers)
        validations = {str(dv.sqref).split(":")[0][0] for dv in wb["TIEU_CHI"].data_validations.dataValidation}
        self.assertTrue({"G", "H", "I", "L"} <= validations)
        self.assertEqual(wb["TIEU_CHI"].cell(2, column(wb["TIEU_CHI"], "trang_thai")).protection.locked, False)
        self.assertEqual(wb["TIEU_CHI"].cell(2, column(wb["TIEU_CHI"], "criterion_id")).protection.locked, True)

    def test_lock_keeps_only_inputs_editable(self):
        path, meta = self.build(lock=True)
        wb = load_workbook(path)
        self.assertTrue(meta["locked"])
        self.assertTrue(wb.security.lockStructure)
        self.assertTrue(all(wb[name].protection.sheet for name in wb.sheetnames))
        self.assertIsNone(wb["TIEU_CHI"].protection.password)

    def test_same_pairs_different_order_per_rater(self):
        a, _ = self.build("A")
        b, _ = self.build("B")
        order = lambda p: [r[1] for r in load_workbook(p)["CAP_CV_JD"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(order(a)), sorted(order(b)))
        again, _ = self.build("A")
        self.assertEqual(order(a), order(again))

    def test_blank_workbook_reports_every_item(self):
        path, meta = self.build()
        issues = ratings_io.check(ratings_io.read(path))
        unrated = [i for i in issues if "chưa chấm" in i.message]
        self.assertEqual(len(unrated), meta["criterion_rows"] + meta["requirement_rows"])

    def test_complete_workbook_exports(self):
        path, _ = self.build()
        fill_complete(path)
        book = ratings_io.read(path)
        self.assertEqual(ratings_io.check(book), [])
        records = ratings_io.to_records(book)
        self.assertEqual([r["pair_id"] for r in records], ["DEMO-P01", "DEMO-P02", "DEMO-P03"])
        first = records[0]
        self.assertEqual(list(first["criteria"]), CRITERIA)
        self.assertEqual(first["criteria"]["C6_certs_activities"], {**first["criteria"]["C6_certs_activities"], "status": "missing_in_cv", "score": 0})
        self.assertEqual(first["criteria"]["C4_projects"]["score"], 3)
        self.assertEqual(first["requirements"]["JD_DEMO_BE-M01"]["status"], "met")
        self.assertEqual(first["overall_fit"], 4)
        self.assertEqual(first["rater_id"], "A")

    def test_inconsistent_entries_are_caught(self):
        path, _ = self.build()
        fill_complete(path)
        wb = load_workbook(path)
        crit, req, overall = wb["TIEU_CHI"], wb["YEU_CAU_JD"], wb["TONG_THE"]
        crit.cell(2, column(crit, "diem")).value = None                  # scored without a score
        crit.cell(3, column(crit, "trich_dan")).value = None             # score 3 without a quote
        crit.cell(4, column(crit, "trang_thai"), "not_applicable")       # keeps its score, no reason
        crit.cell(5, column(crit, "diem"), 7)                            # out of range, pasted past the list
        req.cell(2, column(req, "trich_dan")).value = None               # met without a quote
        overall.cell(2, column(overall, "muc_phu_hop"), 6)
        overall.cell(3, column(overall, "tieu_chi_goi_y_1")).value = None
        wb.save(path)
        messages = [i.message for i in ratings_io.check(ratings_io.read(path))]
        for expected in ("thiếu điểm", "cần trích dẫn", "phải bỏ trống điểm", "cần ghi lý do",
                         "không phải số nguyên 0–4", "met cần trích dẫn", "muc_phu_hop", "mã tiêu chí"):
            self.assertTrue(any(expected in m for m in messages), expected)

    def test_deleted_rows_are_detected(self):
        path, _ = self.build()
        fill_complete(path)
        wb = load_workbook(path)
        wb["TIEU_CHI"].delete_rows(3)
        wb.save(path)
        messages = [i.message for i in ratings_io.check(ratings_io.read(path))]
        self.assertTrue(any("rows were added or deleted" in m for m in messages))
        self.assertTrue(any("đổi thứ tự" in m for m in messages))

    def test_builder_rejects_bad_inputs(self):
        pairs = self.tmp / "pairs.csv"
        pairs.write_text("pair_id,cv_id,jd_id\nP1,CV1,JD_DEMO_BE\nP1,CV2,JD_DEMO_BE\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "duplicate pair_id"):
            self.build(pairs=pairs)
        pairs.write_text("pair_id,cv_id,jd_id\nP1,CV1,JD_MISSING\n", encoding="utf-8")
        with self.assertRaises(FileNotFoundError):
            self.build(pairs=pairs)
        bad_jd = self.tmp / "jd"
        bad_jd.mkdir()
        (bad_jd / "JD_X.json").write_text(json.dumps({"job_title": "x"}), encoding="utf-8")
        pairs.write_text("pair_id,cv_id,jd_id\nP1,CV1,JD_X\n", encoding="utf-8")
        with self.assertRaises(Exception):
            self.build(pairs=pairs, jd=bad_jd)


if __name__ == "__main__":
    unittest.main()
