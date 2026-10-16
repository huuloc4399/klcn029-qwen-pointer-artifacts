"""Inter-rater agreement: statistics against hand-computed values, and the run on two demo workbooks.

Run from the project root: python -m unittest downstream/tests/test_iaa.py -v
"""

from __future__ import annotations

import csv
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from downstream.data import build_rating_workbook as builder
from downstream.data import ratings_io
from downstream.eval import iaa
from downstream.tests.test_rating_workbook import column, fill_complete


class StatisticsTests(unittest.TestCase):
    def test_unweighted_kappa_matches_textbook_table(self):
        # 2x2 table [[20, 5], [10, 15]]: p_o = 0.7, p_e = 0.5, kappa = 0.4.
        a = ["y"] * 25 + ["n"] * 25
        b = ["y"] * 20 + ["n"] * 5 + ["y"] * 10 + ["n"] * 15
        self.assertAlmostEqual(iaa.kappa(a, b, ["y", "n"]), 0.4)

    def test_quadratic_kappa_by_hand(self):
        # Two swaps one level apart over four levels: 1 - (1/18) / (40/144) = 0.8.
        self.assertAlmostEqual(iaa.kappa([1, 2, 3, 4], [1, 2, 4, 3], [1, 2, 3, 4], weighted=True), 0.8)
        self.assertAlmostEqual(iaa.kappa([0, 4], [0, 4], iaa.SCORES, weighted=True), 1.0)

    def test_kappa_undefined_when_both_raters_use_one_level(self):
        self.assertIsNone(iaa.kappa([3, 3, 3], [3, 3, 3], iaa.SCORES, weighted=True))
        self.assertIsNone(iaa.kappa([], [], iaa.SCORES))

    def test_spearman_with_ties(self):
        self.assertAlmostEqual(iaa.spearman([1, 2, 3, 4], [10, 20, 30, 40]), 1.0)
        self.assertAlmostEqual(iaa.spearman([1, 2, 3, 4], [4, 3, 2, 1]), -1.0)
        self.assertAlmostEqual(iaa.spearman([1, 1, 2, 3], [1, 2, 2, 3]), 3.75 / 4.5)  # average ranks for ties
        self.assertIsNone(iaa.spearman([2, 2, 2], [1, 2, 3]))


class AgreementRunTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.paths = {}
        for rater in ("A", "B"):
            path = self.tmp / f"rater_{rater}.xlsx"
            builder.build(builder.DEMO_PAIRS, builder.DEMO_JD_GOLD, rater, path, created_at="2026-10-13T00:00:00Z")
            fill_complete(path)
            self.paths[rater] = path

    def edit(self, rater, sheet, pair, match, values):
        """Set cells on the row of `pair` whose `match` columns equal the given values."""
        wb = load_workbook(self.paths[rater])
        ws = wb[sheet]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, column(ws, "pair_id")).value != pair:
                continue
            if all(ws.cell(r, column(ws, k)).value == v for k, v in match.items()):
                for k, v in values.items():
                    ws.cell(r, column(ws, k)).value = v
                break
        wb.save(self.paths[rater])

    def test_identical_ratings_agree_fully(self):
        report = iaa.run(self.paths["A"], self.paths["B"], None, reps=50)
        m = report["metrics"]
        self.assertEqual((m["pairs"], m["cvs"]), (3, 2))
        self.assertEqual(m["criteria"]["C4_projects"]["exact"], 1.0)
        self.assertIsNone(m["criteria"]["C4_projects"]["qwk"])  # every score is 3
        self.assertEqual(m["criterion_status"]["agreement"], 1.0)
        self.assertEqual(m["requirements"]["agreement"], 1.0)
        self.assertAlmostEqual(m["requirements"]["kappa"], 1.0)
        self.assertEqual(report["disagreement_count"], 0)

    def test_disagreements_and_incomplete_pairs(self):
        self.edit("B", "TIEU_CHI", "DEMO-P01", {"criterion_id": "C4_projects"}, {"diem": 1})
        self.edit("B", "TIEU_CHI", "DEMO-P01", {"criterion_id": "C2_it_content"}, {"diem": 2})  # gap 1: not listed
        self.edit("B", "TIEU_CHI", "DEMO-P03", {"criterion_id": "C7_ats"},
                  {"trang_thai": "unreadable", "diem": None, "ly_do": "PDF hai cột trộn chữ"})
        self.edit("B", "YEU_CAU_JD", "DEMO-P01", {"loai": "must"}, {"trang_thai": "missing", "trich_dan": None})
        self.edit("B", "TONG_THE", "DEMO-P03", {}, {"muc_phu_hop": 2})
        self.edit("A", "TONG_THE", "DEMO-P02", {}, {"muc_phu_hop": None})  # A has not finished P02

        out = self.tmp / "iaa"
        report = iaa.run(self.paths["A"], self.paths["B"], out, reps=100)
        m = report["metrics"]
        self.assertEqual(m["pairs"], 2)
        self.assertEqual(report["rater_a"]["incomplete_pairs"], ["DEMO-P02"])
        c4 = m["criteria"]["C4_projects"]
        self.assertEqual((c4["n"], c4["mae"], c4["gap_ge_2"]), (2, 1.0, 1))
        self.assertEqual(m["criteria"]["C7_ats"]["n"], 1)  # unreadable leaves the score comparison
        self.assertLess(m["criterion_status"]["agreement"], 1.0)
        self.assertEqual(m["requirements"]["confusion"]["met | missing"], 1)

        kinds = sorted(row["kind"] for row in report["disagreement_rows"])
        self.assertEqual(kinds, ["một người chấm được, người kia không", "mức phù hợp lệch 2",
                                 "yêu cầu: met so với missing", "điểm lệch 2"])
        with (out / "disagreements.csv").open(encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0]["decision"], "")
        saved = json.loads((out / "iaa_report.json").read_text(encoding="utf-8"))
        self.assertNotIn("disagreement_rows", saved)  # quotes stay in the CSV only
        text = (out / "IAA_REPORT.md").read_text(encoding="utf-8")
        self.assertIn("Gộp 8 tiêu chí", text)
        self.assertNotIn("Built REST API", text)

    def test_jsonl_export_gives_the_same_numbers(self):
        self.edit("B", "TIEU_CHI", "DEMO-P02", {"criterion_id": "C8_jd_fit"}, {"diem": 1})
        exported = self.tmp / "rater_B.jsonl"
        book = ratings_io.read(self.paths["B"])
        exported.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in ratings_io.to_records(book)),
                            encoding="utf-8")
        from_xlsx = iaa.run(self.paths["A"], self.paths["B"], None, reps=0)["metrics"]
        from_jsonl = iaa.run(self.paths["A"], exported, None, reps=0)["metrics"]
        self.assertEqual(from_xlsx, from_jsonl)

    def test_bootstrap_resamples_cvs(self):
        self.edit("B", "TIEU_CHI", "DEMO-P01", {"criterion_id": "C4_projects"}, {"diem": 1})
        report = iaa.run(self.paths["A"], self.paths["B"], None, reps=200, seed=7)
        ci = report["bootstrap_95ci"]["qwk.pooled"]
        self.assertLessEqual(ci["low"], ci["high"])
        self.assertEqual(report["bootstrap_95ci"], iaa.run(self.paths["A"], self.paths["B"], None, reps=200,
                                                           seed=7)["bootstrap_95ci"])


if __name__ == "__main__":
    unittest.main()
