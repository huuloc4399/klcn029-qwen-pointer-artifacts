"""Rater package: contents, locking and what must never be inside.

Run from the project root: python -m unittest downstream/tests/test_package_rater.py -v
"""

from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from downstream.data import build_rating_workbook as builder
from downstream.data import package_rater


class PackageRaterTests(unittest.TestCase):
    def test_package_holds_workbook_pdfs_and_guide_only(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            fake_pdf = folder / "fake.pdf"
            fake_pdf.write_bytes(b"%PDF-1.4 demo")
            pairs = folder / "pairs.csv"
            pairs.write_text(builder.DEMO_PAIRS.read_text(encoding="utf-8"), encoding="utf-8")
            target = package_rater.package("A", split="dev", pairs=pairs, jd_gold=builder.DEMO_JD_GOLD, jd_text=folder,
                                           out_dir=folder / "out", deadline="20/10/2026", coordinator="Lộc",
                                           pdf_paths=lambda ids: {i: fake_pdf for i in ids})
            self.assertEqual(target.name, "pilot_rater_A.zip")
            with zipfile.ZipFile(target) as archive:
                names = sorted(archive.namelist())
                self.assertEqual(names, ["pilot_rater_A/00_DOC_TRUOC.docx", "pilot_rater_A/CV_PDF/CV_DEMO_01.pdf",
                                         "pilot_rater_A/CV_PDF/CV_DEMO_02.pdf", "pilot_rater_A/pilot_rater_A.xlsx"])
                wb = load_workbook(io.BytesIO(archive.read("pilot_rater_A/pilot_rater_A.xlsx")))
            self.assertTrue(wb.security.lockStructure)
            self.assertNotIn("intended_fit", [c.value for c in wb["CAP_CV_JD"][1]])
            meta = dict(wb["META"].iter_rows(min_row=2, max_col=2, values_only=True))
            self.assertEqual((meta["split"], meta["pair_count"], meta["locked"]), ("dev", 3, True))
            self.assertTrue((folder / "out" / "PACKAGES_LOG.jsonl").is_file())

    def test_guide_only_writes_just_the_guide(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            pairs = folder / "pairs.csv"
            pairs.write_text(builder.DEMO_PAIRS.read_text(encoding="utf-8"), encoding="utf-8")

            def no_pdfs(ids):
                raise AssertionError("guide-only must not touch CV PDFs")

            target = package_rater.package("B", split="dev", pairs=pairs, jd_gold=builder.DEMO_JD_GOLD, jd_text=folder,
                                           out_dir=folder / "out", pdf_paths=no_pdfs, guide_only=True)
            self.assertEqual(target.name, "pilot_rater_B_00_DOC_TRUOC.docx")
            self.assertEqual(sorted(p.name for p in (folder / "out").iterdir()), [target.name])
            text = "\n".join(p.text for p in Document(target).paragraphs)
            self.assertIn(package_rater.GUIDE_VERSION, text)
            for rule in ("Chấm từ dưới lên", "C7 · Khả năng đáp ứng ATS", "Ctrl+A rồi Ctrl+C", "ít nhất 4/5 kỹ năng",
                         "Trước khi chọn missing"):
                self.assertIn(rule, text)
            for rule in package_rater.requirement_rules().RULES:  # word for word what the model prompts hold
                self.assertIn(rule, text)
            self.assertIn("Ghi chú bắt buộc", text)

    def test_missing_pdf_stops_packaging(self):
        with tempfile.TemporaryDirectory() as folder:
            manifest = Path(folder) / "manifest.csv"
            manifest.write_text("id,path,language\nEN001,does/not/exist.pdf,en\n", encoding="utf-8")
            with self.assertRaisesRegex(FileNotFoundError, "EN001"):
                package_rater.cv_pdf_paths({"EN001"}, manifest)


if __name__ == "__main__":
    unittest.main()
